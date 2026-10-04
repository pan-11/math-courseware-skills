"""D1 synthetic groups: actual local files, no media services or real acceptance."""
from copy import deepcopy
from pathlib import Path
import sys
import unittest
import uuid

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills/math-courseware-studio/scripts'))
from runtime import state, workflow, autopilot, review, prompts, image_api, automation_store as store
import test_image_api as api_fixtures
try:
    from runtime import grouped_gates as groups
except ImportError:
    groups = None


class GroupedGateTests(unittest.TestCase):
    def setUp(self):
        self.root = ROOT / 'tests/runs/autopilot-abcd1-20261004/fixtures' / ('d-' + uuid.uuid4().hex)
        state.init_project(self.root, 'Synthetic D1 only', mode='full_course', scope_evidence='Synthetic full course')
        self.data = workflow.load(self.root)

    def ref(self, path):
        return {'path': path, 'sha256': state.sha256(state.resolve(self.root, path))}

    def text(self, path, text='Synthetic actual text, not course acceptance'):
        state.resolve(self.root, path).write_text(text, encoding='utf-8')
        return self.ref(path)

    def image(self, path, color='blue'):
        from PIL import Image
        Image.new('RGB', (100, 100), color).save(state.resolve(self.root, path))
        return self.ref(path)

    def versions(self, *records):
        return {r['path']: r['sha256'] for record in records for r in record['files'].values()}

    def record(self, name, files=None, sources=None, approved=False, **metadata):
        files = files or {'text': self.text('planning/' + name + '.md')}
        sources = sources or {}
        path = 'planning/' + name + '-review.json'
        state.write_json(state.resolve(self.root, path), {'passed': True,
            'source_versions': {**sources, **{r['path']: r['sha256'] for r in files.values()}}})
        result = dict(files=files, source_versions=sources, review=self.ref(path), **metadata)
        if approved:
            path = 'planning/' + name + '-approval.json'
            state.write_json(state.resolve(self.root, path), dict(decision='approved',
                user_evidence='Synthetic fixture approval', targets=list(files.values())))
            result['approval'] = self.ref(path)
        return result

    def save(self):
        state.write_json(state.resolve(self.root, '_state/workflow.json'), self.data)

    def fixture(self, routes=('shots', 'shots'), candidates=True):
        stages = self.data['stages']
        source = self.text('planning/source-review.json', '{"kind":"requirements","basis":"synthetic","limitations":["no real source"]}')
        stages['analysis'] = self.record('analysis', source_review=source)
        self.ids = ['V%03d' % (n + 1) for n in range(len(routes))]
        stages['blueprint'] = self.record('blueprint', approved=True,
            sources=self.versions(stages['analysis']), video_ids=self.ids,
            total_pages=5, coverage=sorted(workflow.BLUEPRINT_COVERAGE))
        images = [self.image('planning/cover-%s.png' % n, color) for n, color in enumerate(('red', 'blue', 'green', 'orange'))]
        assets = [dict(asset_id=name, kind=kind, version='v001', files=[self.image('planning/' + name + '.png')])
                  for name, kind in (('CHAR1', 'character'), ('SCENE1', 'scene'), ('PROP1', 'prop'))]
        state.write_json(state.resolve(self.root, '_state/assets.json'), {'assets': assets})
        files = {'selected': images[0], 'canonical': self.ref('_state/assets.json')}
        files.update({a['asset_id']: a['files'][0] for a in assets})
        choice = dict(kind='candidates' if candidates else 'existing', selected=images[0])
        if candidates: choice['candidates'] = images
        else:
            state.record_approval(self.root, dict(targets=[images[0]], decision='approved', user_evidence='Synthetic prior master'))
        stages['shared-assets'] = self.record('shared', files=files, style_choice=choice)
        self.manifests = {}
        for vid, route in zip(self.ids, routes):
            products = {}
            products['script'] = self.record(vid + '-script', sources=self.versions(stages['blueprint']),
                complete=True, asset_ids=['CHAR1', 'SCENE1', 'PROP1'])
            products['voice'] = self.record(vid + '-voice', sources=self.versions(products['script']), verbatim=True)
            if route == 'shots':
                products['preview'] = self.record(vid + '-preview',
                    files={'image': self.image('planning/' + vid + '-preview.png')},
                    sources=self.versions(products['script']), panel_count=25)
                products['assets'] = self.record(vid + '-assets', files=files,
                    sources=self.versions(stages['shared-assets']))
            else:
                products['first-frame'] = self.record(vid + '-frame',
                    files={'image': self.image('planning/' + vid + '-frame.png')},
                    sources=self.versions(products['script']), color=True)
                products['prompts'] = self.record(vid + '-prompts',
                    sources=self.versions(products['script'], products['first-frame']), complete=True)
            self.manifests[vid] = dict(video_id=vid, route=route, workflow={'products': products})
            if route == 'talking': self.prepare_video(vid)
            self.save_video(vid)
        if 'shots' not in routes: self.prepare_global()
        self.save()

    def save_video(self, vid):
        path = 'planning/' + vid + '-manifest.json'
        state.write_json(state.resolve(self.root, path), self.manifests[vid])
        self.data['videos'][vid] = self.ref(path)

    def prepare_video(self, vid):
        manifest = self.manifests[vid]
        names = ['script', 'voice', 'frame_plan', 'prompts', 'production', 'classroom']
        if manifest['route'] == 'shots': names += ['director', 'board_plan']
        manifest['workflow']['preparation'] = self.record(vid + '-prep',
            files={name: self.text('planning/' + vid + '-prep-' + name + '.md') for name in names},
            sources=self.versions(self.data['stages']['blueprint'], *manifest['workflow']['products'].values()))

    def prepare_global(self):
        self.data['stages']['video-preparation'] = self.record('all-prep', video_ids=self.ids,
            sources=self.versions(self.data['stages']['blueprint'],
                *(m['workflow']['preparation'] for m in self.manifests.values())))

    def directions(self):
        for vid, manifest in self.manifests.items():
            if manifest['route'] != 'shots': continue
            p = manifest['workflow']['products']
            p['director'] = self.record(vid + '-director', sources=self.versions(p['script'], p['voice'], p['assets']))
            p['storyboard'] = self.record(vid + '-board', files={'image': self.image('planning/' + vid + '-board.png')},
                sources=self.versions(p['director'], p['assets']), formal=True, black_white=True)
            p['style'] = self.record(vid + '-style', files={'image': self.image('planning/' + vid + '-style.png')},
                sources=self.versions(p['assets'], p['storyboard']))
            p['prompts'] = self.record(vid + '-prompts', complete=True,
                sources=self.versions(p['script'], p['voice'], p['director'], p['storyboard'], p['style']))
            self.prepare_video(vid)
            self.save_video(vid)
        self.prepare_global()
        self.save()

    def start(self, **extra):
        plan = dict(activation_evidence='Synthetic automatic choice',
            review_policy='grouped_creative_v1', review_policy_evidence='Synthetic explicit grouped choice',
            tasks=[dict(id='analyze', step='analysis', kind='produce', inputs=['planning/analysis.md'],
                outputs=['report'], instruction='Synthetic source review')])
        plan.update(extra)
        return autopilot.start(self.root, plan)

    def snapshot(self, name='video-creative'):
        self.assertIsNotNone(groups, 'D1 needs real group validation')
        return groups.snapshot(self.root, name)

    def approve(self, name='video-creative', evidence='Synthetic whole group approval', **extra):
        snap = self.snapshot(name)
        self.assertEqual(snap['issues'], [])
        return state.record_approval(self.root, dict(targets=snap['targets'], decision='approved',
            user_evidence=evidence, review_group=snap['identity'], **extra))

    def allowed(self, step, **kwargs):
        return workflow.check(self.root, step, **kwargs)

    def test_policy_is_persisted_only_by_explicit_start(self):
        self.fixture()
        result = self.start()
        self.assertEqual(result.get('review_policy'), 'grouped_creative_v1')
        self.assertEqual(result.get('review_policy_evidence'), 'Synthetic explicit grouped choice')

    def test_unknown_policy_rejected(self):
        self.fixture()
        with self.assertRaisesRegex(ValueError, 'review_policy'):
            self.start(review_policy='future_policy')

    def test_policy_needs_its_actual_selection_evidence(self):
        self.fixture()
        with self.assertRaisesRegex(ValueError, 'evidence'):
            self.start(review_policy_evidence='')

    def test_selected_modules_cannot_opt_in(self):
        self.fixture()
        self.data['current_task']['mode'] = 'selected_modules'
        self.save()
        with self.assertRaisesRegex(ValueError, 'full_course'):
            self.start()

    def test_candidates_only_in_explicit_current_automatic_full_course(self):
        self.fixture()
        self.assertFalse(self.allowed('video-assets', video_id='V001')['allowed'])
        self.start()
        self.assertTrue(self.allowed('video-assets', video_id='V001')['allowed'])
        autopilot.control(self.root, mode='manual', evidence='Synthetic manual switch')
        self.assertFalse(self.allowed('video-assets', video_id='V001')['allowed'])
        autopilot.control(self.root, mode='automatic', evidence='Synthetic return')
        self.data['current_task']['evidence'] = 'Different task scope'
        self.save()
        self.assertFalse(self.allowed('video-assets', video_id='V001')['allowed'])

    def test_double_shots_two_groups_reach_pages_without_third_adoption(self):
        self.fixture(); self.start()
        first = self.approve()
        self.assertTrue(self.allowed('video-director', video_id='V001')['allowed'])
        self.directions()
        self.assertFalse(self.allowed('video-upload', video_id='V001')['allowed'])
        self.assertTrue(groups.adopted(self.root, 'video-creative')['approved'])
        second = self.approve('video-direction')
        self.assertNotEqual(first['decision_id'], second['decision_id'])
        self.assertTrue(self.allowed('video-upload', video_id='V002')['allowed'])
        self.assertTrue(self.allowed('pages')['allowed'], self.allowed('pages')['issues'])
        decisions = state.read_lines(state.resolve(self.root, '_state/decisions.jsonl'))
        self.assertEqual(len(decisions), 2)

    def test_mixed_first_group_has_talking_packet_without_future_shots_files(self):
        self.fixture(('shots', 'talking')); self.start()
        snap = self.snapshot()
        self.assertEqual(snap['issues'], [])
        self.assertTrue(any('V002-prep' in r['path'] for r in snap['targets']))
        self.assertFalse(any('V001-prep' in r['path'] for r in snap['targets']))
        self.approve(); self.directions(); self.approve('video-direction')
        self.assertTrue(self.allowed('pages')['allowed'], self.allowed('pages')['issues'])

    def test_all_talking_first_group_covers_global_prep_and_reaches_pages(self):
        self.fixture(('talking', 'talking')); self.start()
        self.approve()
        self.assertTrue(self.allowed('pages')['allowed'], self.allowed('pages')['issues'])
        self.assertTrue(self.allowed('video-upload', video_id='V002')['allowed'])
        self.assertTrue(self.snapshot('video-direction')['issues'])

    def test_individual_approvals_do_not_form_a_group(self):
        self.fixture(); self.start()
        snap = self.snapshot()
        for target in snap['targets']:
            state.record_approval(self.root, dict(targets=[target], user_evidence='Synthetic separate decision'))
        self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])
        self.assertFalse(self.allowed('video-director', video_id='V001')['allowed'])

    def test_partial_targets_cannot_claim_whole_group(self):
        self.fixture(); self.start(); snap = self.snapshot()
        with self.assertRaisesRegex(ValueError, 'exact|complete'):
            state.record_approval(self.root, dict(targets=snap['targets'][:-1],
                user_evidence='Synthetic partial', review_group=snap['identity']))

    def test_no_page_or_image_relaxation_before_group_or_pages(self):
        self.fixture(); self.start()
        self.assertFalse(self.allowed('pages')['allowed'])
        self.assertFalse(self.allowed('page-image')['allowed'])
        self.approve(); self.directions(); self.approve('video-direction')
        self.assertFalse(self.allowed('page-image')['allowed'])

    def image_setup(self, route='builtin'):
        self.fixture()
        for path in ('_state/math.json', '_state/story.json'):
            state.record_approval(self.root, dict(targets=[self.ref(path)], user_evidence='Synthetic adopted teaching core'))
        self.start(preferences={'image': {'route': route, 'evidence': 'Synthetic actual image choice',
            'authorization': {'scope': 'run', 'purposes': ['cover', 'asset'], 'paid_generation': route != 'builtin',
                'include_rework': True, 'scope_evidence': 'Synthetic whole-run consent'}}})

    def prepare_image(self, purpose='asset', target='CHAR1'):
        return prompts.prepare(self.root, {'tasks': [dict(purpose=purpose, target_id=target,
            prompt='Synthetic candidate real preparation; never external image service')]})

    def test_candidate_asset_uses_real_prepare_and_builtin_registration(self):
        self.image_setup()
        batch = self.prepare_image()
        job = state.read_json(state.resolve(self.root, batch['jobs'][0]))
        self.assertIn('creative_candidate', job)
        source = self.image('planning/builtin-result.png')
        result = image_api.register_builtin(self.root, dict(job_path=batch['jobs'][0],
            input_digest=job['input_digest'], image_path=str(state.resolve(self.root, source['path'])),
            tool_evidence='Synthetic local tool response'))
        self.assertEqual(result['status'], 'downloaded')
        self.assertFalse(state.is_approved(self.root, '_state/assets.json'))

    def test_candidate_cover_uses_real_prepare_and_fake_api(self):
        self.image_setup('openai_image_api')
        batch = self.prepare_image('cover', 'COVER01')
        transport = api_fixtures.Transport()
        result = image_api.run_job(self.root, batch['jobs'][0], transport, timeout=1, poll_interval=0)
        self.assertEqual(result['status'], 'downloaded')
        self.assertEqual(len(transport.gets), 1)

    def test_candidate_images_never_waive_teaching_core_approval(self):
        self.fixture(); self.start(preferences={'image': {'route': 'builtin', 'evidence': 'Synthetic choice'}})
        with self.assertRaisesRegex(ValueError, 'confirmation'):
            self.prepare_image()

    def test_candidate_pending_api_rechecks_manual_scope_and_preserves_unknown_recovery(self):
        self.image_setup('openai_image_api')
        batch = self.prepare_image()
        autopilot.control(self.root, mode='manual', evidence='Synthetic stop automatic mode')
        transport = api_fixtures.Transport()
        with self.assertRaises(ValueError):
            image_api.run_job(self.root, batch['jobs'][0], transport, timeout=1, poll_interval=0)
        self.assertEqual(transport.posts, [])

    def independent(self, snapshot):
        paths = [r['path'] for r in snapshot['targets']]
        packet = review.prepare(self.root, dict(rubric='video', producer_id='synthetic-producer',
            artifacts=paths, sources=list(snapshot['source_versions']), instruction='Synthetic full group inspection'))['packet']
        data = state.read_json(state.resolve(self.root, packet))
        review.record(self.root, packet, dict(packet_sha256=state.sha256(state.resolve(self.root, packet)),
            reviewer_id='synthetic-independent', method='independent_agent', source_versions=data['versions'],
            checks=[dict(id=c['id'], verdict='pass', evidence=paths, note='Synthetic inspected artifact') for c in data['criteria']]))
        return packet

    def group_task(self, name='video-creative', packet=None):
        snap = self.snapshot(name)
        paths = [r['path'] for r in snap['targets']]
        return dict(id=name, step=groups.GROUPS[name], kind='human', review_group=name,
            instruction='Display the whole actual group and await one user decision', inputs=paths,
            outputs=['file%02d' % i for i in range(len(paths))], review_packets=[packet] if packet else [])

    def group_receipt(self, action, decision=None):
        group = action.get('review_group')
        self.assertIsInstance(group, dict, 'Dispatch must show concrete current group')
        return dict(task_id=action['task']['id'], claim=action['claim'], status='completed',
            user_evidence='Synthetic actual complete group approval', human_decision='approved',
            artifacts={role: r['path'] for role, r in zip(action['task']['outputs'], group['targets'])},
            review_group=group['identity'], **(dict(decision_id=decision['decision_id'],
                decision_event_id=decision['calibration_event']['event_id']) if decision else {}))

    def test_human_group_dispatch_before_adoption_and_linked_receipt_counts_once(self):
        self.fixture(); self.start()
        snap = self.snapshot(); packet = self.independent(snap)
        autopilot.extend(self.root, [self.group_task(packet=packet)], 'Synthetic actual group queue')
        action = autopilot.next_task(self.root, 'host')['human_tasks'][0]
        receipt = self.group_receipt(action)
        with self.assertRaisesRegex(ValueError, 'decision'):
            autopilot.record(self.root, receipt)
        decision = self.approve(evidence=receipt['user_evidence'], review_packet=packet)
        from runtime import calibration
        before = len(calibration.snapshots(self.root, store.load_run(self.root)))
        self.assertEqual(autopilot.record(self.root, self.group_receipt(action, decision))['status'], 'done')
        saved = store.load_run(self.root)
        self.assertEqual(saved['tasks'][-1]['last_human_receipt']['event_id'], decision['calibration_event']['event_id'])
        self.assertEqual(len(calibration.snapshots(self.root, store.load_run(self.root))), before)

    def test_group_requires_independent_review_to_dispatch_but_pass_does_not_adopt(self):
        self.fixture(); self.start()
        autopilot.extend(self.root, [self.group_task()], 'Synthetic group task')
        action = autopilot.next_task(self.root, 'host')
        self.assertEqual(action['human_tasks'], [])
        self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])

    def test_generic_old_decision_cannot_replay_as_group_authority(self):
        self.fixture(); self.start(); snap = self.snapshot()
        old = state.record_approval(self.root, dict(targets=snap['targets'], user_evidence='Same synthetic words'))
        grouped = self.approve(evidence='Same synthetic words')
        self.assertNotEqual(old['calibration_event']['event_id'], grouped['calibration_event']['event_id'])
        self.assertEqual(grouped['review_group'], snap['identity'])

    def test_direction_adoption_requires_every_current_shots_video(self):
        self.fixture(); self.start(); self.approve(); self.directions()
        self.manifests['V002']['workflow']['products'].pop('storyboard')
        self.save_video('V002'); self.save()
        self.assertTrue(self.snapshot('video-direction')['issues'])

    def test_current_asset_file_cannot_be_replaced_with_approved_copy(self):
        self.fixture(); self.start(); self.approve()
        path = state.resolve(self.root, '_state/assets.json')
        changed = state.read_json(path); changed['assets'][0]['version'] = 'v002'; state.write_json(path, changed)
        self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])

    def test_changed_source_route_pointer_or_group_member_invalidates_adoption(self):
        for mutation in ('source', 'route', 'pointer', 'member'):
            with self.subTest(mutation=mutation):
                self.setUp(); self.fixture(); self.start(); self.approve()
                if mutation == 'source':
                    self.text('planning/blueprint.md', 'Synthetic revised blueprint')
                elif mutation == 'route':
                    self.manifests['V001']['route'] = 'talking'; self.save_video('V001'); self.save()
                elif mutation == 'pointer':
                    path = 'planning/new-manifest.json'
                    state.write_json(state.resolve(self.root, path), self.manifests['V001'])
                    self.data['videos']['V001'] = self.ref(path); self.save()
                else:
                    self.text('planning/V001-script.md', 'Synthetic revised script')
                self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])

    def test_rejected_group_member_cannot_pass_through_old_group(self):
        self.fixture(); self.start(); decision = self.approve()
        state.record_approval(self.root, dict(targets=[decision['targets'][0]], decision='rejected',
            user_evidence='Synthetic later rejection'))
        self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])

    def test_new_policy_rejects_separate_H19_and_individual_H06(self):
        self.fixture()
        for node in ('H19', 'H06'):
            with self.subTest(node=node), self.assertRaisesRegex(ValueError, 'H19|H06'):
                self.start(tasks=[dict(id='separate', step='analysis', kind='human', human_node=node,
                    inputs=['planning/analysis.md'], outputs=['report'], instruction='Synthetic old gate')])

    def final_fixture(self):
        from pptx import Presentation
        self.fixture(('talking',)); self.start(); self.approve()
        stages = self.data['stages']
        picture = stages['shared-assets']['files']['selected']
        pages = [dict(page_id='P%02d' % i, order=i, video_ids=['V001'] if i == 2 else [],
            native_objects=[{'object_id': 'video-frame-V001'}] if i == 2 else [], image=picture)
            for i in range(1, 6)]
        state.write_json(state.resolve(self.root, '_state/pages.json'), {'pages': pages})
        stages['pages'] = self.record('pages', files={'pages': self.ref('_state/pages.json')}, approved=True,
            sources=self.versions(stages['blueprint'], stages['video-preparation']))
        stages['images'] = self.record('images', files={p['page_id']: picture for p in pages}, approved=True,
            sources=self.versions(stages['pages'], stages['shared-assets']))
        deck = Presentation()
        for _ in pages: deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(state.resolve(self.root, 'planning/editable.pptx'))
        stages['editable'] = self.record('editable',
            files={'pptx': self.ref('planning/editable.pptx'), 'limitations': self.text('planning/limitations.md')},
            sources=self.versions(stages['pages'], stages['images']))
        stages['documents'] = self.record('documents', files={role: self.text('planning/' + role + '.md')
            for role in ('classroom-script', 'lesson-presentation', 'lesson-plan', 'worksheet', 'blackboard')},
            sources=self.versions(stages['pages'], stages['images']))
        path = state.resolve(self.root, 'planning/actual-media.mp4')
        path.write_bytes(b'\x00\x00\x00\x18ftypisom\x00\x00\x02\x00isomiso2')
        self.manifests['V001']['workflow']['final'] = self.record('final-V001',
            files={'media': self.ref('planning/actual-media.mp4')}, approved=True)
        self.save_video('V001'); self.save()

    def finish_reports(self):
        stages = self.data['stages']
        final = self.manifests['V001']['workflow']['final']
        media = final['files']['media']
        path = 'planning/playback.json'
        state.write_json(state.resolve(self.root, path), dict(passed=True, inspection='user_report',
            basis='SYNTHETIC report only; no actual playback occurred',
            source_versions={**self.versions(stages['editable']), media['path']: media['sha256']}))
        self.manifests['V001']['workflow']['final'] = self.record('final-V001-checked',
            files={'media': media, 'playback': self.ref(path)}, approved=True)
        self.save_video('V001')
        versions = self.versions(stages['pages'], stages['editable'], stages['documents'])
        versions[media['path']] = media['sha256']
        report = dict(passed=True, inspection='user_report', basis='SYNTHETIC H23 only; no WPS inspection',
            source_versions=versions, checks={key: True for key in groups.FINAL_CHECKS})
        state.write_json(state.resolve(self.root, 'planning/wps.json'), report)
        stages['delivery'] = self.record('delivery', files={'manifest': self.text('planning/delivery.md'),
            'wps': self.ref('planning/wps.json')}, sources=versions)
        self.save()

    def test_H23_dispatch_shows_uninspected_real_files_then_completion_requires_checks(self):
        self.final_fixture()
        autopilot.extend(self.root, [dict(id='acceptance', kind='human', step='complete', human_node='H23',
            inputs=['planning/editable.pptx'], outputs=['inspection'],
            instruction='Actually inspect editable objects, every page, all playback and supporting materials')],
            'Synthetic final acceptance task')
        action = autopilot.next_task(self.root, 'host')
        human = next(t for t in action['human_tasks'] if t['task']['id'] == 'acceptance')
        self.assertEqual(set(human.get('final_checklist', [])), set(groups.FINAL_CHECKS))
        self.assertFalse(self.allowed('complete')['allowed'])
        result = dict(task_id='acceptance', claim=human['claim'], status='completed',
            user_evidence='Synthetic actual test report', artifacts={'inspection': 'planning/limitations.md'})
        with self.assertRaises(ValueError): autopilot.record(self.root, result)
        self.finish_reports()
        self.assertTrue(self.allowed('complete')['allowed'], self.allowed('complete')['issues'])
        result['artifacts']['inspection'] = 'planning/wps.json'
        self.assertEqual(autopilot.record(self.root, result)['status'], 'done')
        self.assertFalse(autopilot.status(self.root)['whole_course_complete'])

    def test_H23_missing_object_or_usage_checks_blocks_complete(self):
        self.final_fixture(); self.finish_reports()
        report = state.read_json(state.resolve(self.root, 'planning/wps.json'))
        report['checks'].pop('teaching_graphics')
        state.write_json(state.resolve(self.root, 'planning/wps.json'), report)
        files = self.data['stages']['delivery']['files']; files['wps'] = self.ref('planning/wps.json')
        self.data['stages']['delivery'] = self.record('rechecked-delivery', files=files,
            sources=self.data['stages']['delivery']['source_versions'])
        self.save()
        self.assertFalse(self.allowed('complete')['allowed'])

    def test_H23_media_text_cannot_pass_actual_container_gate(self):
        self.final_fixture(); self.finish_reports()
        self.text('planning/actual-media.mp4', 'a video script only')
        self.assertFalse(self.allowed('complete')['allowed'])

    def test_four_folder_group_paths_match_legacy_behavior(self):
        self.root = ROOT / 'tests/runs/autopilot-abcd1-20261004/fixtures' / ('d-four-' + uuid.uuid4().hex)
        state.init_project(self.root, 'Synthetic D1 four folders', mode='full_course',
            scope_evidence='Synthetic full course', layout='four-folders')
        self.data = workflow.load(self.root)
        self.fixture(('talking',)); self.start(); self.approve()
        self.assertTrue(self.allowed('pages')['allowed'], self.allowed('pages')['issues'])
        self.assertTrue((self.root / '02_work/_state/decisions.jsonl').is_file())

    def test_old_run_without_policy_uses_original_assets_gate(self):
        self.fixture()
        autopilot.start(self.root, dict(activation_evidence='Synthetic original automatic mode',
            tasks=[dict(id='source', step='analysis', inputs=['planning/analysis.md'],
                outputs=['report'], instruction='Synthetic source review')]))
        self.assertNotIn('review_policy', store.load_run(self.root))
        self.assertFalse(self.allowed('video-assets', video_id='V001')['allowed'])

    def test_existing_adopted_master_reuses_without_four_new_candidates(self):
        self.fixture(('talking',), candidates=False); self.start(); self.approve()
        self.assertTrue(self.allowed('pages')['allowed'])

    def test_missing_role_preview_voice_or_unselected_cover_blocks_group(self):
        for missing in ('role', 'preview', 'voice', 'selected'):
            with self.subTest(missing=missing):
                self.setUp(); self.fixture(); self.start()
                if missing == 'role':
                    self.data['stages']['shared-assets']['files'].pop('PROP1')
                elif missing == 'selected':
                    self.data['stages']['shared-assets']['style_choice']['selected'] = self.image('planning/not-a-candidate.png')
                else:
                    product = self.manifests['V002']['workflow']['products'][missing]
                    product['panel_count' if missing == 'preview' else 'verbatim'] = 0
                    self.save_video('V002')
                self.save()
                self.assertTrue(self.snapshot()['issues'])

    def test_missing_canonical_record_returns_group_diagnostics(self):
        self.fixture(); self.start()
        # Retain the bytes in place: a missing logical canonical source is modeled
        # through a reader failure, not deletion of any fixture or user file.
        from unittest.mock import patch
        original = state.sha256
        def unavailable(path):
            if Path(path) == state.resolve(self.root, '_state/assets.json'): raise FileNotFoundError('Synthetic missing canonical')
            return original(path)
        with patch.object(state, 'sha256', unavailable):
            self.assertTrue(self.snapshot()['issues'])

    def test_candidate_reference_rejection_and_changed_basis_block_before_requests(self):
        for mode in ('reference-rejected', 'canonical-changed', 'scope-changed'):
            with self.subTest(mode=mode):
                self.setUp(); self.image_setup('openai_image_api')
                batch = prompts.prepare(self.root, {'tasks': [dict(purpose='asset', target_id='CHAR1',
                    prompt='Synthetic referenced candidate', reference_assets=[{'asset_id': 'PROP1', 'version': 'v001'}])]})
                if mode == 'reference-rejected':
                    target = self.data['stages']['shared-assets']['files']['PROP1']
                    state.record_approval(self.root, dict(targets=[target], decision='rejected', user_evidence='Synthetic rejected prop'))
                elif mode == 'canonical-changed':
                    self.text('_state/assets.json', '{"assets":[]}')
                else:
                    self.data['current_task']['mode'] = 'selected_modules'; self.save()
                transport = api_fixtures.Transport()
                with self.assertRaises(ValueError):
                    image_api.run_job(self.root, batch['jobs'][0], transport, timeout=1, poll_interval=0)
                self.assertEqual(transport.posts, [])

    def test_candidate_unknown_submission_is_not_retried_after_manual_switch(self):
        self.image_setup('openai_image_api'); batch = self.prepare_image()
        transport = api_fixtures.Transport(TimeoutError('Synthetic uncertain provider'))
        first = image_api.run_job(self.root, batch['jobs'][0], transport, timeout=1, poll_interval=0)
        self.assertEqual(first['status'], 'submission_unknown')
        autopilot.control(self.root, mode='manual', evidence='Synthetic stop')
        again = image_api.run_job(self.root, batch['jobs'][0], transport, resume=True, timeout=1, poll_interval=0)
        self.assertEqual(again['status'], 'submission_unknown')
        self.assertEqual(len(transport.posts), 1)
        self.assertEqual(autopilot.status(self.root)['image_budget_status']['used'], 1)

    def test_candidate_builtin_rechecks_basis_before_registering_current_result(self):
        self.image_setup(); batch = self.prepare_image()
        job = state.read_json(state.resolve(self.root, batch['jobs'][0]))
        source = self.image('planning/late-image.png')
        autopilot.control(self.root, mode='manual', evidence='Synthetic mode switch')
        with self.assertRaises(ValueError):
            image_api.register_builtin(self.root, dict(job_path=batch['jobs'][0], input_digest=job['input_digest'],
                image_path=str(state.resolve(self.root, source['path'])), tool_evidence='Synthetic result'))
        self.assertEqual(state.read_json(state.resolve(self.root, batch['jobs'][0]))['status'], 'pending')

    def test_group_rejected_receipt_retains_waiting_and_one_occurrence(self):
        self.fixture(); self.start(); snap = self.snapshot(); packet = self.independent(snap)
        autopilot.extend(self.root, [self.group_task(packet=packet)], 'Synthetic queue extension')
        action = autopilot.next_task(self.root, 'host')['human_tasks'][0]
        decision = state.record_approval(self.root, dict(targets=snap['targets'], decision='rejected',
            user_evidence='Synthetic user rejects the actual group', review_group=snap['identity'], review_packet=packet))
        receipt = self.group_receipt(action, decision)
        receipt.update(status='returned', human_decision='rejected', user_evidence=decision['user_evidence'])
        self.assertEqual(autopilot.record(self.root, receipt)['status'], 'waiting_external')
        self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])
        self.assertEqual(store.load_run(self.root)['tasks'][-1]['last_human_receipt']['event_id'], decision['calibration_event']['event_id'])

    def test_non_group_receipt_cannot_acquire_group_scope_by_annotation(self):
        self.fixture(); self.start()
        task = dict(id='plain', kind='human', step='analysis', inputs=['planning/analysis.md'],
            outputs=['report'], instruction='Synthetic ordinary response')
        autopilot.extend(self.root, [task], 'Synthetic ordinary task')
        action = autopilot.next_task(self.root, 'host')['human_tasks'][0]
        with self.assertRaisesRegex(ValueError, 'group'):
            autopilot.record(self.root, dict(task_id='plain', claim=action['claim'], status='completed',
                artifacts={'report': 'planning/analysis.md'}, user_evidence='Synthetic reply',
                review_group=self.snapshot()['identity']))

    def test_group_adoption_in_manual_does_not_gain_automatic_policy_authority(self):
        self.fixture(); self.start(); snap = self.snapshot()
        autopilot.control(self.root, mode='manual', evidence='Synthetic manual')
        with self.assertRaises(ValueError):
            state.record_approval(self.root, dict(targets=snap['targets'], review_group=snap['identity'],
                user_evidence='Synthetic group words outside automatic policy'))

    def test_scope_reconciliation_does_not_reuse_old_group_identity(self):
        self.fixture(); self.start(); self.approve()
        self.data['current_task']['evidence'] = 'Synthetic newly reconciled scope'; self.save()
        autopilot.reconcile(self.root, evidence='Synthetic actual scope instruction')
        self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])

    def test_existing_master_keeps_valid_legacy_approval_file(self):
        self.fixture(('talking',)); self.start()
        shared = self.data['stages']['shared-assets']
        shared['style_choice'] = {'kind': 'existing', 'selected': shared['style_choice']['selected']}
        path = 'planning/old-master-adoption.json'
        state.write_json(state.resolve(self.root, path), dict(decision='approved',
            user_evidence='Synthetic actual preserved master decision', targets=[shared['style_choice']['selected']]))
        shared['approval'] = self.ref(path); self.save()
        self.approve()
        self.assertTrue(self.allowed('pages')['allowed'])

    def test_independent_group_review_must_cover_unselected_candidates_as_sources(self):
        self.fixture(); self.start(); snap = self.snapshot()
        paths = [r['path'] for r in snap['targets']]
        packet = review.prepare(self.root, dict(rubric='video', producer_id='synthetic-maker',
            artifacts=paths, sources=paths, instruction='Synthetic incomplete source coverage'))['packet']
        data = state.read_json(state.resolve(self.root, packet))
        review.record(self.root, packet, dict(packet_sha256=state.sha256(state.resolve(self.root, packet)),
            reviewer_id='synthetic-reviewer', method='independent_agent', source_versions=data['versions'],
            checks=[dict(id=c['id'], verdict='pass', evidence=paths, note='Synthetic review omits candidates') for c in data['criteria']]))
        autopilot.extend(self.root, [self.group_task(packet=packet)], 'Synthetic incomplete review group')
        self.assertEqual(autopilot.next_task(self.root, 'host')['human_tasks'], [])

    def test_H23_missing_current_supporting_file_keeps_final_gate_closed(self):
        self.final_fixture(); self.finish_reports()
        self.text('planning/worksheet.md', 'Synthetic later worksheet revision')
        self.assertFalse(self.allowed('complete')['allowed'])

    def test_actual_later_readoption_of_rejected_same_group_is_new_event(self):
        self.fixture(); self.start(); first = self.approve()
        current = self.snapshot()
        rejected = dict(targets=current['targets'], review_group=current['identity'], decision='rejected',
            user_evidence='Synthetic explicit rejection of shown group')
        refusal = state.record_approval(self.root, rejected)
        self.assertEqual(state.record_approval(self.root, rejected)['calibration_event']['event_id'], refusal['calibration_event']['event_id'])
        self.assertFalse(groups.adopted(self.root, 'video-creative')['approved'])
        changed_mind = state.record_approval(self.root, dict(targets=current['targets'],
            review_group=current['identity'], user_evidence='Synthetic genuine later readoption of these exact files'))
        self.assertNotEqual(first['calibration_event']['event_id'], changed_mind['calibration_event']['event_id'])
        self.assertTrue(groups.adopted(self.root, 'video-creative')['approved'])

    def test_member_generic_readoption_cannot_revive_prior_whole_group(self):
        from unittest.mock import patch
        for name in ('video-creative', 'video-direction'):
            with self.subTest(group=name):
                self.setUp(); self.fixture(); self.start()
                if name == 'video-direction':
                    self.approve(); self.directions()
                with patch.object(state, 'now', return_value='2026-10-05T12:00:00Z'):
                    first = self.approve(name)
                current = self.snapshot(name)
                product = 'script' if name == 'video-creative' else 'director'
                target = next(iter(self.manifests['V001']['workflow']['products'][product]['files'].values()))
                with patch.object(state, 'now', return_value='2026-10-05T11:00:00Z'):
                    refusal = state.record_approval(self.root, dict(targets=[target], decision='rejected',
                        user_evidence='Synthetic later member rejection despite reversed wall clock'))
                    member = state.record_approval(self.root, dict(targets=[target], decision='approved',
                        user_evidence='Synthetic readoption of this member only'))
                self.assertTrue(state.is_approved(self.root, target['path']))
                self.assertGreater(refusal['calibration_event']['event_seq'], first['calibration_event']['event_seq'])
                self.assertFalse(groups.adopted(self.root, name)['approved'])
                step = 'video-director' if name == 'video-creative' else 'video-upload'
                self.assertFalse(self.allowed(step, video_id='V001')['allowed'])
                opinion = dict(targets=current['targets'], review_group=current['identity'], decision='approved',
                    user_evidence=first['user_evidence'])
                renewed = state.record_approval(self.root, opinion)
                self.assertGreater(renewed['calibration_event']['event_seq'], member['calibration_event']['event_seq'])
                self.assertNotEqual(first['calibration_event']['event_id'], renewed['calibration_event']['event_id'])
                self.assertTrue(groups.adopted(self.root, name)['approved'])
                self.assertTrue(self.allowed(step, video_id='V001')['allowed'])
                self.assertEqual(state.record_approval(self.root, opinion)['calibration_event']['event_id'],
                                 renewed['calibration_event']['event_id'])

    def test_H23_dispatch_includes_real_delivery_files_before_inspection_metadata(self):
        self.final_fixture()
        interactive = self.text('planning/selected-interactive.html', '<p>Synthetic classroom interaction</p>')
        self.data['stages']['delivery'] = {'files': {'interactive': interactive,
            'wps': {'path': 'planning/future-wps.json', 'sha256': '0' * 64},
            'manifest': {'path': 'planning/future-delivery.md', 'sha256': '0' * 64}}}
        self.save()
        autopilot.extend(self.root, [dict(id='acceptance', kind='human', step='complete', human_node='H23',
            inputs=['planning/editable.pptx'], outputs=['inspection'],
            instruction='Inspect all real delivery artifacts including the selected interaction')],
            'Synthetic final inspection task before its report')
        human = next(t for t in autopilot.next_task(self.root, 'host')['human_tasks']
                     if t['task']['id'] == 'acceptance')
        self.assertEqual(human['input_versions'].get(interactive['path']), interactive['sha256'])
        self.assertNotIn('planning/future-wps.json', human['input_versions'])
        self.assertNotIn('planning/future-delivery.md', human['input_versions'])
        self.assertFalse(self.allowed('complete')['allowed'])
        self.data['stages']['delivery']['files']['interactive'] = {'path': 'planning/missing-interactive.html',
                                                                  'sha256': '1' * 64}
        self.save()
        self.assertFalse(groups.final_dispatch(self.root)['allowed'])

    def test_H23_inspection_must_bind_actual_delivery_supports_without_self_dependency(self):
        self.final_fixture(); self.finish_reports()
        interactive = self.text('planning/selected-interactive.html', '<p>Synthetic classroom interaction</p>')
        delivery = self.data['stages']['delivery']
        delivery['files']['interactive'] = interactive
        self.data['stages']['delivery'] = self.record('delivery-with-interaction', files=delivery['files'],
            sources=delivery['source_versions'])
        self.save()
        self.assertFalse(self.allowed('complete')['allowed'])
        report = state.read_json(state.resolve(self.root, 'planning/wps.json'))
        report['source_versions'][interactive['path']] = interactive['sha256']
        state.write_json(state.resolve(self.root, 'planning/wps.json'), report)
        files = self.data['stages']['delivery']['files']
        files['wps'] = self.ref('planning/wps.json')
        self.data['stages']['delivery'] = self.record('delivery-all-inspected', files=files,
            sources=report['source_versions'])
        self.save()
        self.assertTrue(self.allowed('complete')['allowed'], self.allowed('complete')['issues'])
        dispatched = groups.final_dispatch(self.root)['inspection_versions']
        self.assertEqual(dispatched.get(interactive['path']), interactive['sha256'])
        self.assertNotIn(files['wps']['path'], dispatched)
        self.assertNotIn(files['manifest']['path'], dispatched)
        self.text(interactive['path'], '<p>Synthetic later changed interaction</p>')
        self.assertFalse(groups.final_dispatch(self.root)['allowed'])
        self.assertFalse(self.allowed('complete')['allowed'])

    def test_grouped_prepare_reuses_pre_run_cover_asset_pending_and_submitted_jobs(self):
        from unittest.mock import patch
        from runtime import image_budget
        for purpose, target in (('cover', 'COVER01'), ('asset', 'CHAR1')):
            for status in ('pending', 'running', 'submission_unknown'):
                with self.subTest(purpose=purpose, status=status):
                    self.setUp(); self.fixture()
                    project = state.load_project(self.root); project['image_route'] = 'openai_image_api'
                    state.write_json(state.resolve(self.root, '_state/project.json'), project)
                    state.record_approval(self.root, dict(targets=[self.ref(p) for p in
                        ('_state/math.json', '_state/story.json', '_state/assets.json')],
                        user_evidence='Synthetic existing adopted definitions before a run'))
                    original = self.prepare_image(purpose, target)
                    path = state.resolve(self.root, original['jobs'][0])
                    job = state.read_json(path)
                    if status != 'pending':
                        job.update(status=status, task_id='provider-test-id')
                        state.write_json(path, job)
                    before = path.read_bytes()
                    payload = image_api.build_payload(self.root, job)
                    self.start(preferences={'image': dict(route='openai_image_api', evidence='Synthetic same route',
                        authorization=dict(scope='run', purposes=['cover', 'asset'], paid_generation=True,
                            include_rework=True, scope_evidence='Synthetic actual current consent'))},
                        limits={'max_images': 1, 'evidence': 'Synthetic one generation limit'})
                    reused = self.prepare_image(purpose, target)
                    self.assertEqual(reused['jobs'], original['jobs'])
                    self.assertEqual(path.read_bytes(), before)
                    self.assertEqual(image_api.build_payload(self.root, state.read_json(path)), payload)
                    transport = api_fixtures.Transport()
                    with patch.object(image_api, 'HttpTransport', lambda key: transport):
                        result = image_api.run_batch(self.root, reused, 'synthetic-key', resume=status != 'pending')
                    self.assertEqual(result['results'][0]['status'],
                                     'submission_unknown' if status == 'submission_unknown' else 'downloaded')
                    self.assertEqual([p[0] for p in transport.posts],
                        [image_api.ENDPOINT, image_api.RESULT_ENDPOINT] if status == 'pending' else
                        [image_api.RESULT_ENDPOINT] if status == 'running' else [])
                    self.assertEqual(image_budget.summary(store.load_run(self.root))['used'], int(status == 'pending'))
                    after = path.read_bytes()
                    self.assertEqual(self.prepare_image(purpose, target)['jobs'], original['jobs'])
                    self.assertEqual(path.read_bytes(), after)

    def test_candidate_prepare_after_same_file_adoption_keeps_original_job(self):
        from runtime import image_budget
        for purpose, target in (('cover', 'COVER01'), ('asset', 'CHAR1')):
            with self.subTest(purpose=purpose):
                self.setUp(); self.image_setup('openai_image_api')
                batch = self.prepare_image(purpose, target)
                path = state.resolve(self.root, batch['jobs'][0]); before = path.read_bytes()
                old = state.read_json(path)
                self.assertIn('creative_candidate', old)
                state.record_approval(self.root, dict(targets=[self.ref('_state/assets.json')],
                    user_evidence='Synthetic later adoption of unchanged canonical assets'))
                self.assertEqual(self.prepare_image(purpose, target)['jobs'], batch['jobs'])
                self.assertEqual(path.read_bytes(), before)
                transport = api_fixtures.Transport()
                self.assertEqual(image_api.run_job(self.root, batch['jobs'][0], transport,
                    timeout=1, poll_interval=0)['status'], 'downloaded')
                self.assertEqual(image_budget.summary(store.load_run(self.root))['used'], 1)
                after = path.read_bytes()
                self.assertEqual(self.prepare_image(purpose, target)['jobs'], batch['jobs'])
                self.assertEqual(path.read_bytes(), after)

    def test_reused_candidate_keeps_original_scope_and_blueprint_restrictions(self):
        from runtime import image_budget
        for change in ('scope', 'blueprint', 'manual'):
            with self.subTest(change=change):
                self.setUp(); self.image_setup('openai_image_api')
                batch = self.prepare_image()
                path = state.resolve(self.root, batch['jobs'][0]); before = path.read_bytes()
                state.record_approval(self.root, dict(targets=[self.ref('_state/assets.json')],
                    user_evidence='Synthetic same-file adoption does not erase candidate scope'))
                if change == 'scope':
                    self.data['current_task']['mode'] = 'selected_modules'
                    self.data['current_task']['modules'] = ['plan', 'video']
                    self.save()
                elif change == 'manual':
                    autopilot.control(self.root, mode='manual', evidence='Synthetic explicit manual switch')
                else:
                    old = self.data['stages']['blueprint']
                    self.data['stages']['blueprint'] = self.record('new-blueprint', approved=True,
                        sources=old['source_versions'], video_ids=self.ids, total_pages=5,
                        coverage=sorted(workflow.BLUEPRINT_COVERAGE))
                    self.save()
                with self.assertRaises(ValueError): self.prepare_image()
                transport = api_fixtures.Transport()
                with self.assertRaises(ValueError):
                    image_api.run_job(self.root, batch['jobs'][0], transport, timeout=1, poll_interval=0)
                self.assertEqual(transport.posts, [])
                self.assertEqual(path.read_bytes(), before)
                self.assertEqual(image_budget.summary(store.load_run(self.root))['used'], 0)

    def quality_human(self, kind):
        if kind == 'group':
            self.fixture(); self.start()
            return self.group_task(packet=self.independent(self.snapshot()))
        self.final_fixture()
        return dict(id='acceptance', kind='human', step='complete', human_node='H23',
            inputs=['planning/editable.pptx'], outputs=['inspection'], instruction='Synthetic actual final inspection')

    def test_group_and_H23_respect_explicit_preference_dependencies(self):
        for kind in ('group', 'H23'):
            with self.subTest(kind=kind):
                self.setUp(); task = self.quality_human(kind)
                group = 'video' if kind == 'group' else 'editable'
                task['requires_preferences'] = [group]
                autopilot.extend(self.root, [task], 'Synthetic explicit human dependency')
                action = autopilot.next_task(self.root, 'host')
                self.assertNotIn(task['id'], [t['task']['id'] for t in action['human_tasks']])
                saved = next(t for t in store.load_run(self.root)['tasks'] if t['spec']['id'] == task['id'])
                self.assertEqual(saved['status'], 'pending')
                self.assertTrue(any('Missing preference: ' + group + '.' in issue for issue in saved['issues']))
                settings = dict(platform='Synthetic platform', model='Synthetic model', sound='Synthetic sound') if kind == 'group' else dict(route='A', entry='full')
                autopilot.configure(self.root, {group: dict(**settings, evidence='Synthetic actual choices')})
                action = autopilot.next_task(self.root, 'host')
                self.assertIn(task['id'], [t['task']['id'] for t in action['human_tasks']])

    def test_group_and_H23_enforce_declared_image_budget_and_authorization(self):
        from runtime import image_budget
        for kind in ('group', 'H23'):
            with self.subTest(kind=kind):
                self.setUp(); task = self.quality_human(kind)
                task.update(requires_preferences=['image'], image_requests=[
                    dict(purpose='cover', target_id='COVER%02d' % n, version='v001') for n in (1, 2)])
                autopilot.configure(self.root, {'limits': {'max_images': 1, 'evidence': 'Synthetic one-image limit'},
                    'image': dict(route='openai_image_api', evidence='Synthetic selected route',
                        authorization=dict(scope='run', purposes=['cover'], paid_generation=True,
                            include_rework=False, scope_evidence='Synthetic consent'))})
                autopilot.extend(self.root, [task], 'Synthetic declared image work on this human task')
                action = autopilot.next_task(self.root, 'host')
                self.assertNotIn(task['id'], [t['task']['id'] for t in action['human_tasks']])
                saved = next(t for t in store.load_run(self.root)['tasks'] if t['spec']['id'] == task['id'])
                self.assertTrue(any('Image budget limit' in issue for issue in saved['issues']))
                autopilot.configure(self.root, {'limits': {'max_images': 2, 'evidence': 'Synthetic actual increase'},
                    'image': {'authorization': {'paid_generation': False, 'scope_evidence': 'Synthetic consent withdrawn'}}})
                action = autopilot.next_task(self.root, 'host')
                self.assertNotIn(task['id'], [t['task']['id'] for t in action['human_tasks']])
                saved = next(t for t in store.load_run(self.root)['tasks'] if t['spec']['id'] == task['id'])
                self.assertTrue(any('paid generation authorization' in issue for issue in saved['issues']))
                autopilot.configure(self.root, {'image': {'authorization': {
                    'paid_generation': True, 'scope_evidence': 'Synthetic renewed actual consent'}}})
                action = autopilot.next_task(self.root, 'host')
                self.assertIn(task['id'], [t['task']['id'] for t in action['human_tasks']])
                self.assertEqual(image_budget.summary(store.load_run(self.root))['used'], 0)

    def test_full_image_budget_keeps_ordinary_group_adoption_available(self):
        from runtime import image_budget
        self.image_setup('openai_image_api')
        autopilot.configure(self.root, {'limits': {'max_images': 1, 'evidence': 'Synthetic one-image limit'}})
        batch = self.prepare_image()
        image_api.run_job(self.root, batch['jobs'][0], api_fixtures.Transport(), timeout=1, poll_interval=0)
        self.assertEqual(image_budget.summary(store.load_run(self.root))['remaining'], 0)
        packet = self.independent(self.snapshot())
        autopilot.extend(self.root, [self.group_task(packet=packet)], 'Synthetic ordinary adoption without new generation')
        action = autopilot.next_task(self.root, 'host')
        human = next(t for t in action['human_tasks'] if t['task']['id'] == 'video-creative')
        decision = self.approve(review_packet=packet)
        self.assertEqual(autopilot.record(self.root, self.group_receipt(human, decision))['status'], 'done')
        self.assertEqual(image_budget.summary(store.load_run(self.root))['used'], 1)


if __name__ == '__main__':
    unittest.main()
