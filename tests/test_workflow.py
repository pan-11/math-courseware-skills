"""Workflow dependency tests with retained synthetic artifacts and no media services."""
import importlib.util
import json
from pathlib import Path
import sys
import unittest
import uuid

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills/math-courseware-studio/scripts'))
from runtime import state
from workflow_fixture import evidence, ref, enable_modules, versions
try:
    from runtime import workflow
except ImportError:
    workflow = None


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(workflow, 'A shared workflow checker is required')
        self.root = ROOT / 'tests/runs' / ('workflow-' + uuid.uuid4().hex)
        state.init_project(self.root, 'Synthetic workflow only')
        self.data = {'schema_version': '1.0', 'project_mode': 'full_course',
                     'scope_evidence': 'Synthetic request to rebuild a full course',
                     'current_task': {'mode': 'full_course', 'modules': ['studio'],
                                      'evidence': 'Synthetic continue instruction'},
                     'stages': {}, 'inputs': {}, 'videos': {}}
        self.save()

    def save(self):
        state.write_json(self.root / '_state/workflow.json', self.data)

    def allowed(self, step, **kwargs):
        self.save()
        return workflow.check(self.root, step, **kwargs)['allowed']

    def source_analysis(self):
        review = {'kind': 'uploaded_courseware', 'page_count': 3,
                  'page_map': [{'page': n, 'role': 'Synthetic lesson function'} for n in range(1, 4)],
                  'overall_quality': 'Clear task sequence, weak closing feedback.',
                  'strengths': ['P1 motivates measurement'], 'improvements': ['P3 needs assessment'],
                  'teaching_flow': ['Context', 'Inquiry', 'Application'],
                  'video_inventory': {'complete': True, 'unique_count': 1, 'occurrence_count': 2,
                      'items': [{'video_id': 'original-1', 'pages': [1, 3], 'inspection': 'viewed',
                                 'content': 'Two children compare measurement units.',
                                 'function': 'Introduce need for a common unit.',
                                 'basis': 'Synthetic media inspected; not a real lesson'}], 'limitations': []}}
        state.write_json(self.root / 'planning/source-review.json', review)
        record = evidence(self.root, 'analysis')
        record['source_review'] = ref(self.root, 'planning/source-review.json')
        self.data['stages']['analysis'] = record
        return review

    def blueprint(self, videos=('V001',)):
        self.source_analysis()
        self.data['stages']['blueprint'] = evidence(self.root, 'blueprint', approved=True,
            sources=versions(self.data['stages']['analysis']),
            video_ids=list(videos), total_pages=5,
            coverage=['teaching_flow', 'activities', 'assessment', 'story_math',
                      'timing', 'page_budget', 'video_inventory', 'source_mapping'])

    def video(self, vid='V001', route='shots', products=(), prepare=False, final=False):
        path = 'planning/' + vid + '-manifest.json'
        checks = {'products': {name: evidence(self.root, vid + '-' + name, approved=True)
                               for name in products}}
        for name in set(products) & {'preview', 'first-frame', 'storyboard', 'assets'}:
            from PIL import Image
            image_path = 'planning/' + vid + '-' + name + '.png'
            Image.new('RGB', (32, 18), 'blue').save(self.root / image_path)
            checks['products'][name] = evidence(self.root, vid + '-' + name,
                files={'image': ref(self.root, image_path)}, approved=True)
        chains = {'preview': ('script',), 'voice': ('script',),
                  'director': ('script', 'voice', 'assets'), 'storyboard': ('director', 'assets'),
                  'style': ('assets', 'storyboard'), 'first-frame': ('script',),
                  'prompts': ('script', 'first-frame') if route == 'talking'
                             else ('script', 'voice', 'director', 'storyboard', 'style')}
        for name, product in list(checks['products'].items()):
            deps = [checks['products'][dep] for dep in chains.get(name, ()) if dep in checks['products']]
            if name == 'script' and self.data['current_task']['mode'] == 'full_course':
                deps.append(self.data['stages'].get('blueprint', {}))
            checks['products'][name] = evidence(self.root, vid + '-' + name,
                files=product['files'], approved=True, sources=versions(*deps))
        if prepare:
            names = ['script', 'voice', 'frame_plan', 'prompts', 'production', 'classroom']
            if route == 'shots':
                names += ['director', 'board_plan']
            files = {}
            for name in names:
                item = evidence(self.root, vid + '-prep-' + name)
                files[name] = item['files']['report']
            checks['preparation'] = evidence(self.root, vid + '-prep', files=files, approved=True,
                sources=versions(self.data['stages'].get('blueprint', {}), *checks['products'].values()),
                missing_assets=[{'asset_id': 'char', 'reason': 'Synthetic tool unavailable',
                                 'acquire': 'Generate from saved prompt', 'blocks': ['page-image']}])
        if final:
            checks['final'] = evidence(self.root, vid + '-final', approved=True)
        state.write_json(self.root / path, {'video_id': vid, 'route': route, 'workflow': checks})
        self.data['videos'][vid] = ref(self.root, path)

    def prep(self, videos=('V001',)):
        self.blueprint(videos)
        for vid in videos:
            self.video(vid, prepare=True)
        self.data['stages']['video-preparation'] = evidence(self.root, 'video-preparation', approved=True,
            video_ids=list(videos), sources=versions(self.data['stages']['blueprint'],
                *(state.read_json(self.root / self.data['videos'][vid]['path'])['workflow']['preparation'] for vid in videos)))

    def shared_assets(self):
        from PIL import Image
        path = 'planning/adopted-style.png'
        Image.new('RGB', (32, 18), 'blue').save(self.root / path)
        selected = ref(self.root, path)
        self.data['stages']['shared-assets'] = evidence(self.root, 'shared-assets', approved=True,
            files={'style': selected}, style_choice={'kind': 'existing', 'selected': selected})

    def test_story_selection_does_not_replace_blueprint(self):
        self.source_analysis()
        self.data['stages']['story'] = evidence(self.root, 'story', approved=True)
        self.assertFalse(self.allowed('video-script', video_id='V001'))

    def test_full_analysis_needed_before_blueprint(self):
        self.assertFalse(self.allowed('blueprint'))
        self.source_analysis()
        self.assertTrue(self.allowed('blueprint'))

    def test_analysis_has_visible_quality_page_flow_and_video_inventory(self):
        source = self.source_analysis()
        source['page_map'] = source['page_map'][:2]
        state.write_json(self.root / 'planning/source-review.json', source)
        self.data['stages']['analysis']['source_review'] = ref(self.root, 'planning/source-review.json')
        self.assertFalse(self.allowed('blueprint'))

    def test_video_counts_deduplicate_media_and_preserve_occurrences(self):
        source = self.source_analysis()
        source['video_inventory']['unique_count'] = 2
        state.write_json(self.root / 'planning/source-review.json', source)
        self.data['stages']['analysis']['source_review'] = ref(self.root, 'planning/source-review.json')
        self.assertFalse(self.allowed('blueprint'))

    def test_unknown_external_video_is_honest_analysis_not_fake_zero(self):
        source = self.source_analysis()
        inventory = source['video_inventory']
        inventory.update(complete=False, unique_count=None, limitations=['External clip unavailable'])
        inventory['items'][0].update(inspection='unavailable', content='Unknown actual content')
        state.write_json(self.root / 'planning/source-review.json', source)
        self.data['stages']['analysis']['source_review'] = ref(self.root, 'planning/source-review.json')
        self.assertTrue(self.allowed('blueprint'))
        inventory['limitations'] = []
        state.write_json(self.root / 'planning/source-review.json', source)
        self.data['stages']['analysis']['source_review'] = ref(self.root, 'planning/source-review.json')
        self.assertFalse(self.allowed('blueprint'))

    def test_unreviewed_blueprint_does_not_allow_video(self):
        self.blueprint()
        self.data['stages']['blueprint'].pop('approval')
        self.assertFalse(self.allowed('video-script', video_id='V001'))

    def test_every_planned_video_required_before_pages(self):
        self.prep(('V001', 'V002'))
        self.assertTrue(self.allowed('pages'))
        self.data['videos'].pop('V002')
        self.assertFalse(self.allowed('pages'))

    def test_concrete_prep_can_continue_without_real_video_but_not_final(self):
        self.prep()
        self.assertTrue(self.allowed('pages'))
        self.assertFalse(self.allowed('complete'))
        self.assertFalse(self.allowed('page-image'))

    def test_video_focus_keeps_full_course_target(self):
        self.data['current_task']['modules'] = ['video']
        self.assertFalse(self.allowed('video-script', video_id='V001'))
        self.assertEqual(workflow.summary(self.root)['project_mode'], 'full_course')

    def test_scoped_refill_accepts_real_external_inputs_without_story(self):
        from fixture_factory import layered_deck
        deck = layered_deck(self.root)
        self.data = enable_modules(self.root, ['editable'], deck=deck.relative_to(self.root).as_posix(), route='B')
        self.assertTrue(self.allowed('editable-build'))
        self.assertFalse(state.is_approved(self.root, '_state/story.json'))
        self.assertFalse(self.allowed('editable-handoff', route='A'))

    def test_refill_missing_text_or_deck_is_specific(self):
        self.data = enable_modules(self.root, ['editable'], route='B')
        self.assertFalse(self.allowed('editable-build'))
        result = workflow.check(self.root, 'editable-build')
        self.assertTrue(any('deck' in item for item in result['issues']))

    def test_scoped_course_maintenance_preserves_project_goal(self):
        from fixture_factory import layered_deck
        deck = layered_deck(self.root)
        self.data = enable_modules(self.root, ['editable'], deck=deck.relative_to(self.root).as_posix())
        self.data['project_mode'] = 'full_course'
        self.assertTrue(self.allowed('editable-build'))
        self.assertEqual(workflow.summary(self.root)['project_mode'], 'full_course')
        self.assertFalse(self.allowed('complete'))

    def test_source_version_change_blocks_dependent_step(self):
        self.source_analysis()
        path = self.root / self.data['stages']['analysis']['files']['report']['path']
        path.write_text('Changed analysis', encoding='utf-8')
        self.assertFalse(self.allowed('blueprint'))

    def test_unclassified_diagnosis_allowed_but_production_denied(self):
        self.data['project_mode'] = 'unclassified'
        self.assertFalse(self.allowed('pages'))
        self.assertEqual(workflow.summary(self.root)['project_mode'], 'unclassified')

    def test_module_selection_does_not_allow_unrequested_work(self):
        self.data = enable_modules(self.root, ['documents'])
        self.assertTrue(self.allowed('documents'))
        self.assertFalse(self.allowed('pages'))

    def test_standalone_video_needs_no_course_project(self):
        root = ROOT / 'tests/runs' / ('standalone-' + uuid.uuid4().hex)
        root.mkdir()
        (root / 'AGENTS.md').write_text('Synthetic standalone workflow. Retain all files.', encoding='utf-8')
        (root / '_state/qa').mkdir(parents=True)
        enable_modules(root, ['video'])
        self.assertTrue(workflow.check(root, 'video-script', video_id='V001')['allowed'])
        self.assertFalse((root / '_state/project.json').exists())

    def test_talking_materials_continue_without_intermediate_adoption(self):
        self.data = enable_modules(self.root, ['video'])
        self.video(route='talking', products=['script', 'first-frame', 'prompts'])
        path = self.root / self.data['videos']['V001']['path']
        manifest = state.read_json(path)
        for product in manifest['workflow']['products'].values():
            product.pop('approval', None)
        for route in ('talking', 'fixed_talking'):
            manifest['route'] = route
            state.write_json(path, manifest)
            self.data['videos']['V001'] = ref(self.root, path.relative_to(self.root).as_posix())
            for step in ('video-assets', 'video-prompts', 'video-upload'):
                with self.subTest(route=route, step=step):
                    self.assertTrue(self.allowed(step, video_id='V001'))
            self.assertEqual(state.read_json(path), manifest)

    def test_talking_direct_delivery_keeps_image_and_review_checks(self):
        self.data = enable_modules(self.root, ['video'])
        self.video(route='talking', products=['script', 'first-frame'])
        path = self.root / self.data['videos']['V001']['path']
        manifest = state.read_json(path)
        products = manifest['workflow']['products']
        for product in products.values():
            product.pop('approval', None)
        state.write_json(path, manifest)
        self.data['videos']['V001'] = ref(self.root, path.relative_to(self.root).as_posix())
        self.assertTrue(self.allowed('video-prompts', video_id='V001'))
        products['first-frame'] = evidence(self.root, 'text-only-frame',
            sources=versions(products['script']))
        state.write_json(path, manifest)
        self.data['videos']['V001'] = ref(self.root, path.relative_to(self.root).as_posix())
        self.assertFalse(self.allowed('video-prompts', video_id='V001'))

    def test_talking_direct_delivery_does_not_reuse_rejected_frame(self):
        self.data = enable_modules(self.root, ['video'])
        self.video(route='talking', products=['script', 'first-frame'])
        path = self.root / self.data['videos']['V001']['path']
        manifest = state.read_json(path)
        products = manifest['workflow']['products']
        for product in products.values():
            product.pop('approval', None)
        state.write_json(path, manifest)
        self.data['videos']['V001'] = ref(self.root, path.relative_to(self.root).as_posix())
        self.assertTrue(self.allowed('video-prompts', video_id='V001'))
        state.record_approval(self.root, {
            'targets': list(products['first-frame']['files'].values()),
            'decision': 'rejected', 'user_evidence': 'Synthetic rejected robot identity'})
        self.assertFalse(self.allowed('video-prompts', video_id='V001'))

    def test_shot_assets_still_need_script_adoption(self):
        self.data = enable_modules(self.root, ['video'])
        self.video(products=['script', 'preview'])
        self.assertTrue(self.allowed('video-assets', video_id='V001'))
        path = self.root / self.data['videos']['V001']['path']
        manifest = state.read_json(path)
        manifest['workflow']['products']['script'].pop('approval')
        state.write_json(path, manifest)
        self.data['videos']['V001'] = ref(self.root, path.relative_to(self.root).as_posix())
        self.assertFalse(self.allowed('video-assets', video_id='V001'))

    def test_talking_route_does_not_require_director_or_board(self):
        self.data = enable_modules(self.root, ['video'])
        self.video(route='talking', products=['script', 'first-frame'])
        self.assertTrue(self.allowed('video-prompts', video_id='V001'))
        self.assertFalse(self.allowed('video-director', video_id='V001'))

    def test_upload_map_revision_does_not_require_preview(self):
        self.data = enable_modules(self.root, ['video-prompts'])
        self.data['inputs']['video'] = self.data['inputs']['video-prompts']
        self.video(products=['script', 'voice', 'director', 'storyboard', 'style', 'prompts'])
        self.assertTrue(self.allowed('video-upload', video_id='V001'))

    def test_draft_cannot_self_declare_not_applicable(self):
        self.data['stages']['analysis'] = {'status': 'not_applicable'}
        self.assertFalse(self.allowed('blueprint'))

    def test_unrecognized_step_fails_closed(self):
        with self.assertRaises(ValueError):
            workflow.require(self.root, 'skip-to-finish')

    def test_new_shot_assets_require_actual_preview(self):
        self.data = enable_modules(self.root, ['video'])
        self.video(products=['script'])
        self.assertFalse(self.allowed('video-assets', video_id='V001'))
        self.video(products=['script', 'preview'])
        self.assertTrue(self.allowed('video-assets', video_id='V001'))

    def test_text_asset_list_cannot_replace_actual_images(self):
        self.data = enable_modules(self.root, ['video'])
        self.video(products=['script', 'voice', 'assets'])
        path = self.root / self.data['videos']['V001']['path']
        manifest = state.read_json(path)
        manifest['workflow']['products']['assets'] = evidence(self.root, 'text-assets', approved=True)
        state.write_json(path, manifest)
        self.data['videos']['V001'] = ref(self.root, path.relative_to(self.root).as_posix())
        self.assertFalse(self.allowed('video-director', video_id='V001'))

    def test_page_plan_must_reserve_all_videos_inside_budget(self):
        self.prep(('V001', 'V002'))
        pages = [{'page_id': 'P' + str(n), 'order': n, 'video_ids': [], 'native_objects': []}
                 for n in range(1, 6)]
        pages[1].update(video_ids=['V001'], native_objects=[{'object_id': 'video-frame-V001'}])
        state.write_json(self.root / '_state/pages.json', {'pages': pages})
        self.data['stages']['pages'] = evidence(self.root, 'pages', approved=True,
            files={'pages': ref(self.root, '_state/pages.json')})
        self.data['stages']['shared-assets'] = evidence(self.root, 'shared-assets', approved=True)
        self.assertFalse(self.allowed('page-image'))

    def test_final_claim_needs_all_actual_media_and_playback(self):
        self.prep()
        for name in ('pages', 'shared-assets', 'images', 'editable', 'documents', 'delivery'):
            self.data['stages'][name] = evidence(self.root, name, approved=True)
        self.assertFalse(self.allowed('complete'))

    def test_empty_scope_cannot_hide_course_goal_in_manifest(self):
        self.data['current_task'] = {'mode': 'selected_modules', 'modules': ['editable'], 'evidence': ''}
        self.assertFalse(self.allowed('editable-build'))

    def test_malformed_page_map_returns_diagnostic(self):
        source = self.source_analysis()
        source['page_map'].append('invalid row')
        state.write_json(self.root / 'planning/source-review.json', source)
        self.data['stages']['analysis']['source_review'] = ref(self.root, 'planning/source-review.json')
        self.assertFalse(self.allowed('blueprint'))

    def test_rejected_current_version_cannot_reuse_old_approval_file(self):
        self.blueprint()
        files = list(self.data['stages']['blueprint']['files'].values())
        state.record_approval(self.root, {'targets': files, 'decision': 'rejected',
            'user_evidence': 'Synthetic rejection after earlier approval'})
        self.assertFalse(self.allowed('video-script', video_id='V001'))

    def test_canonical_page_mapping_cannot_be_replaced_by_an_unrelated_copy(self):
        self.prep()
        page_file = 'planning/other-pages.json'
        state.write_json(self.root / page_file, {'pages': [
            {'page_id': str(i), 'order': i, 'video_ids': ['V001'] if i == 2 else [],
             'native_objects': [{'object_id': 'video-frame-V001'}] if i == 2 else []}
            for i in range(1, 6)]})
        self.data['stages']['pages'] = evidence(self.root, 'wrong-page-plan', approved=True,
            files={'pages': ref(self.root, page_file)})
        self.data['stages']['shared-assets'] = evidence(self.root, 'shared-assets', approved=True)
        self.assertFalse(self.allowed('page-image'))

    def test_valid_full_page_plan_can_proceed_and_phase_exports_are_not_completion(self):
        self.prep()
        path = '_state/pages.json'
        state.write_json(self.root / path, {'pages': [
            {'page_id': 'P' + str(i), 'order': i, 'video_ids': ['V001'] if i == 2 else [],
             'native_objects': [{'object_id': 'video-frame-V001'}] if i == 2 else []}
            for i in range(1, 6)]})
        self.data['stages']['pages'] = evidence(self.root, 'page-plan', approved=True,
            sources=versions(self.data['stages']['blueprint'], self.data['stages']['video-preparation']),
            files={'pages': ref(self.root, path)})
        self.shared_assets()
        self.assertTrue(self.allowed('page-image'))
        self.assertTrue(self.allowed('collect'))
        self.assertFalse(self.allowed('complete'))

    def test_cover_style_cannot_be_skipped_or_only_named(self):
        self.test_valid_full_page_plan_can_proceed_and_phase_exports_are_not_completion()
        self.data['stages']['shared-assets'] = evidence(self.root, 'shared-assets', approved=True)
        self.assertFalse(self.allowed('page-image'))

    def test_new_cover_choice_requires_four_real_candidates(self):
        self.test_valid_full_page_plan_can_proceed_and_phase_exports_are_not_completion()
        choice = self.data['stages']['shared-assets']['style_choice']
        choice.update(kind='candidates', candidates=[choice['selected']])
        self.assertFalse(self.allowed('page-image'))

    def test_independent_document_delivery_does_not_need_course_assets(self):
        self.data = enable_modules(self.root, ['documents'])
        self.data['stages']['delivery'] = evidence(self.root, 'document-delivery')
        self.assertTrue(self.allowed('complete'))

    def test_retained_old_blueprint_does_not_keep_preparation_current(self):
        self.prep()
        self.assertTrue(self.allowed('pages'))
        old = self.data['stages']['blueprint']
        self.data['stages']['blueprint'] = evidence(self.root, 'blueprint-v2', approved=True,
            video_ids=old['video_ids'], total_pages=old['total_pages'], coverage=old['coverage'],
            sources=old['source_versions'])
        self.assertFalse(self.allowed('pages'))

    def test_retained_old_director_does_not_keep_upload_current(self):
        self.data = enable_modules(self.root, ['video-prompts'])
        self.video(products=['director', 'prompts'])
        self.assertTrue(self.allowed('video-upload', video_id='V001'))
        path = self.data['videos']['V001']['path']
        manifest = state.read_json(self.root / path)
        manifest['workflow']['products']['director'] = evidence(self.root, 'director-v2', approved=True)
        state.write_json(self.root / path, manifest)
        self.data['videos']['V001'] = ref(self.root, path)
        self.assertFalse(self.allowed('video-upload', video_id='V001'))

    def test_text_only_deliverables_cannot_complete_course(self):
        self.test_valid_full_page_plan_can_proceed_and_phase_exports_are_not_completion()
        for name in ('images', 'editable', 'documents', 'delivery'):
            self.data['stages'][name] = evidence(self.root, name, approved=True)
        path = self.data['videos']['V001']['path']
        manifest = state.read_json(self.root / path)
        text = evidence(self.root, 'fake-media')['files']['report']
        manifest['workflow']['final'] = evidence(self.root, 'fake-final', approved=True,
            files={'media': text, 'playback': text})
        state.write_json(self.root / path, manifest)
        self.data['videos']['V001'] = ref(self.root, path)
        self.assertFalse(self.allowed('complete'))

    def test_corrupt_image_returns_diagnostic_instead_of_crashing(self):
        import base64
        self.data = enable_modules(self.root, ['video'])
        self.video(route='talking', products=['script', 'first-frame'])
        path = self.data['videos']['V001']['path']
        manifest = state.read_json(self.root / path)
        image_path = 'planning/bad-crc.png'
        raw = bytearray(base64.b64decode(
            'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4//8/AAX+Av4N70a4AAAAAElFTkSuQmCC'))
        raw[45] ^= 1
        (self.root / image_path).write_bytes(raw)
        manifest['workflow']['products']['first-frame'] = evidence(self.root, 'bad-image', approved=True,
            files={'image': ref(self.root, image_path)})
        state.write_json(self.root / path, manifest)
        self.data['videos']['V001'] = ref(self.root, path)
        self.assertFalse(self.allowed('video-prompts', video_id='V001'))
        self.assertTrue(any('invalid image bytes' in issue for issue in
            workflow.check(self.root, 'video-prompts', video_id='V001')['issues']))

    def test_selected_style_must_itself_be_an_image(self):
        self.test_valid_full_page_plan_can_proceed_and_phase_exports_are_not_completion()
        image = self.data['stages']['shared-assets']['files']['style']
        text = evidence(self.root, 'style-description')['files']['report']
        self.data['stages']['shared-assets'] = evidence(self.root, 'mixed-style', approved=True,
            files={'unrelated-image': image, 'style': text}, style_choice={'kind': 'existing', 'selected': text})
        self.assertFalse(self.allowed('page-image'))
        self.assertTrue(any('selected-style: actual images required' in issue for issue in
            workflow.check(self.root, 'page-image')['issues']))

    def completed_files(self):
        """A synthetic evidence-graph fixture, not media generation or real playback acceptance."""
        from pptx import Presentation
        self.test_valid_full_page_plan_can_proceed_and_phase_exports_are_not_completion()
        stages = self.data['stages']
        page_data = state.read_json(self.root / '_state/pages.json')
        image = stages['shared-assets']['files']['style']
        for page in page_data['pages']:
            page['image'] = image
        state.write_json(self.root / '_state/pages.json', page_data)
        stages['pages'] = evidence(self.root, 'final-pages', approved=True,
            files={'pages': ref(self.root, '_state/pages.json')},
            sources=versions(stages['blueprint'], stages['video-preparation']))
        stages['images'] = evidence(self.root, 'final-images', approved=True,
            files={p['page_id']: image for p in page_data['pages']}, sources=versions(stages['pages'], stages['shared-assets']))
        deck = Presentation()
        for _ in page_data['pages']:
            deck.slides.add_slide(deck.slide_layouts[6])
        path = 'planning/synthetic-output.pptx'
        deck.save(self.root / path)
        stages['editable'] = evidence(self.root, 'final-editable', files={'pptx': ref(self.root, path)},
            sources=versions(stages['pages'], stages['images']))
        documents = {role: evidence(self.root, 'doc-' + role)['files']['report']
                     for role in ('classroom-script', 'lesson-presentation', 'lesson-plan')}
        stages['documents'] = evidence(self.root, 'final-documents', files=documents,
            sources=versions(stages['pages'], stages['images']))
        # Fixed container marker fixture tests format classification only; it contains no encoded movie.
        media_path = 'planning/container-marker.mp4'
        (self.root / media_path).write_bytes(b'\x00\x00\x00\x18ftypisom\x00\x00\x02\x00isomiso2')
        media = ref(self.root, media_path)
        playback_path = 'planning/synthetic-playback.json'
        state.write_json(self.root / playback_path, {'passed': True, 'inspection': 'user_report',
            'basis': 'SYNTHETIC evidence only; no real user report or playback occurred.',
            'source_versions': {**versions(stages['editable']), media['path']: media['sha256']}})
        path = self.data['videos']['V001']['path']
        manifest = state.read_json(self.root / path)
        manifest['workflow']['final'] = evidence(self.root, 'actual-final-fixture', approved=True,
            files={'media': media, 'playback': ref(self.root, playback_path)})
        state.write_json(self.root / path, manifest)
        self.data['videos']['V001'] = ref(self.root, path)
        sources = versions(stages['editable'], stages['documents'])
        wps = 'planning/synthetic-wps.json'
        state.write_json(self.root / wps, {'passed': True, 'inspection': 'user_report',
            'basis': 'SYNTHETIC evidence only; no WPS inspection occurred.', 'source_versions': sources})
        stages['delivery'] = evidence(self.root, 'final-delivery', sources=sources,
            files={'manifest': evidence(self.root, 'delivery-list')['files']['report'], 'wps': ref(self.root, wps)})

    def test_complete_accepts_current_artifacts_then_rejects_renamed_text_video(self):
        self.completed_files()
        self.assertTrue(self.allowed('complete'), workflow.check(self.root, 'complete')['issues'])
        (self.root / 'planning/container-marker.mp4').write_text('Only a video production plan', encoding='utf-8')
        self.assertFalse(self.allowed('complete'))
        path = self.data['videos']['V001']['path']
        manifest = state.read_json(self.root / path)
        manifest['workflow']['final']['files']['media'] = ref(self.root, 'planning/container-marker.mp4')
        state.write_json(self.root / path, manifest)
        self.data['videos']['V001'] = ref(self.root, path)
        self.save()
        self.assertTrue(any('video container required' in issue for issue in
            workflow.check(self.root, 'complete')['issues']))

    def test_complete_rejects_text_disguised_as_pptx_and_wrong_page_count(self):
        from pptx import Presentation
        self.completed_files()
        path = self.root / 'planning/synthetic-output.pptx'
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(path)
        self.assertFalse(self.allowed('complete'))
        # Refresh only the inspected file pointer, so this assertion checks the format/page gate itself.
        stages = self.data['stages']
        stages['editable']['files']['pptx'] = ref(self.root, 'planning/synthetic-output.pptx')
        self.save()
        self.assertTrue(any('full page count required' in issue for issue in workflow.check(self.root, 'complete')['issues']))
        path.write_text('Editable PPT plan only', encoding='utf-8')
        stages['editable']['files']['pptx'] = ref(self.root, 'planning/synthetic-output.pptx')
        self.save()
        self.assertTrue(any('cannot read PPTX' in issue for issue in workflow.check(self.root, 'complete')['issues']))

    def test_board_adoption_must_cover_grouped_director_without_extra_approval(self):
        self.data = enable_modules(self.root, ['video-prompts'])
        self.video(products=['script', 'voice', 'director', 'storyboard', 'style'])
        path = self.data['videos']['V001']['path']
        manifest = state.read_json(self.root / path)
        products = manifest['workflow']['products']
        products['director'].pop('approval')
        state.write_json(self.root / path, manifest)
        self.data['videos']['V001'] = ref(self.root, path)
        self.assertFalse(self.allowed('video-prompts', video_id='V001'))
        approval_path = products['storyboard']['approval']['path']
        decision = state.read_json(self.root / approval_path)
        decision['targets'].extend(products['director']['files'].values())
        state.write_json(self.root / approval_path, decision)
        products['storyboard']['approval'] = ref(self.root, approval_path)
        state.write_json(self.root / path, manifest)
        self.data['videos']['V001'] = ref(self.root, path)
        self.assertTrue(self.allowed('video-prompts', video_id='V001'))

    def test_independent_B_background_repair_stays_in_editable_scope(self):
        self.data = enable_modules(self.root, ['editable'], route='B')
        self.save()
        self.assertTrue(workflow.require_image(self.root, 'repair')['allowed'])
        self.data['route_choice']['route'] = 'A'
        self.save()
        with self.assertRaisesRegex(ValueError, 'outside the requested modules'):
            workflow.require_image(self.root, 'repair')


if __name__ == '__main__':
    unittest.main()
