"""Retained synthetic layered PPTX fixture; never a real lesson acceptance."""
from pathlib import Path
from datetime import datetime
import uuid
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE


def retained_project(label="pptx"):
    root = Path(__file__).parent / "runs" / (label + "-" + datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:8])
    root.mkdir(parents=True)
    (root / "AGENTS.md").write_text("# Synthetic fixture rules\ninputs: source fixtures; editable: returned copies, outputs, retained work; _state: synthetic review evidence. No deletion. Not real courseware acceptance.\n", encoding="utf-8")
    for name in ("inputs", "editable", "_state"):
        (root / name).mkdir()
    return root


def layered_deck(root):
    root = Path(root)
    picture = root / "inputs" / "fixture-image.png"
    image = Image.new("RGB", (400, 300), "#b6dce8")
    ImageDraw.Draw(image).rectangle((50, 40, 330, 230), fill="#4f9780")
    image.save(picture)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(960), Pt(540)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    background = slide.shapes.add_picture(str(picture), Pt(0), Pt(0), Pt(960), Pt(540))
    background.name = "background-image"
    image_shape = slide.shapes.add_picture(str(picture), Pt(500), Pt(220), Pt(300), Pt(225))
    image_shape.name = "cropped-picture"
    image_shape.crop_left, image_shape.crop_right = .1, .15
    group = slide.shapes.add_group_shape()
    group.name = "math-group"
    for i in range(2):
        child = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(100 + 100 * i), Pt(260), Pt(80), Pt(80))
        child.name = "math-block-" + str(i + 1)
    title = slide.shapes.add_textbox(Pt(60), Pt(45), Pt(750), Pt(70))
    title.name = "existing-title"
    title.text = "原有文字 3 + 2 = 5"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    source = root / "inputs" / "layered-source.pptx"
    deck.save(source)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("customXml/retained.xml", '<retained xmlns="urn:fixture">unknown part retained</retained>')
    return source
