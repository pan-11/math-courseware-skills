"""Retained synthetic export checks, not real-course or WPS visual acceptance."""
from datetime import datetime
from pathlib import Path
import sys
import hashlib
import html
import unittest
import uuid
import zipfile

from PIL import Image
from pptx import Presentation
from docx import Document
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1] /
                       'skills/math-courseware-studio/scripts'))
from runtime import state, exports
from workflow_fixture import enable_modules


def fixture(register_scope=True):
    root = Path(__file__).parent.resolve() / 'runs' / (
        'exports-' + datetime.now().strftime('%Y%m%d-%H%M%S') + '-' + uuid.uuid4().hex[:8])
    state.init_project(root, '合成数学课件测试，不是真实首课')
    pages = []
    for number, size in [(2, (400, 300)), (1, (800, 450))]:
        path = root / 'slides' / f'fixture-{number}.png'
        Image.new('RGB', size, 'red' if number == 1 else 'blue').save(path)
        pages.append({'page_id': f'page-{number}', 'order': number,
                      'image': {'path': path.relative_to(root).as_posix(),
                                'sha256': state.sha256(path)}})
    state.write_json(root / '_state/pages.json', {'pages': pages})
    paths = list(exports.CORE) + [p['image']['path'] for p in pages]
    state.record_approval(root, {'user_evidence': 'Synthetic fixture approval only; not user course approval',
                                'targets': [{'path': p, 'sha256': state.sha256(root / p)} for p in paths]})
    versions = {p: state.sha256(root / p) for p in ['_state/pages.json', *[page['image']['path'] for page in pages]]}
    common = {'source_versions': versions, 'page_order': ['page-1', 'page-2']}
    documents = {
        'classroom-script': {'title': '合成测试课堂逐字稿', 'content': [
            {'page_id': f'page-{n}', 'title': f'第{n}个数学任务',
             'paragraphs': [f'同学们，请观察第{n}组图形，比较大小。', '数学符号：3 < 5 & 8 > 6。']}
            for n in (1, 2)]},
        'lesson-presentation': {'title': '合成测试公开课说课稿', 'content': [
            {'heading': '教材分析', 'paragraphs': ['通过观察与表达建立数学关系。'],
             'page_ids': ['page-1', 'page-2']}]},
        'lesson-plan': {'title': '合成测试教学设计', 'content': {
            'metadata': [{'heading': '教学目标', 'paragraphs': ['学生能比较数量并说出依据。']}],
            'rows': [{'page_ids': ['page-1'], 'cells': ['导入', '教师出示图形。', '学生观察。',
                                                      '能指出不同。', '激活已有经验。']},
                     {'page_ids': ['page-2'], 'cells': ['探究', '\n'.join(
                         f'第{i:03d}步：引导学生比较并解释，保留符号 3 < 5 & 8 > 6。' for i in range(90)),
                         '学生逐项记录比较结果。', '根据表达检查推理依据。', '课后填写观察与改进。']}],
            'afterword': [{'heading': '板书设计', 'paragraphs': ['观察数量→比较→说明依据。']}]}}
    }
    for kind, record in documents.items():
        state.write_json(root / f'_state/documents/{kind}.json', {**common, **record})
    if register_scope:
        enable_modules(root)
    return root


def textless_selection(root):
    pages = sorted(state.read_json(root / '_state/pages.json')['pages'], key=lambda p: p['order'])
    selected = []
    for number, page in enumerate(pages, 1):
        target = root / 'slides' / f'textless-{number}.png'
        with Image.open(root / page['image']['path']) as image:
            size = image.size
        Image.new('RGB', size, 'green' if number == 1 else 'yellow').save(target)
        selected.append({'page_id': page['page_id'],
                         'image': {'path': target.relative_to(root).as_posix(), 'sha256': state.sha256(target)}})
    state.record_approval(root, {'user_evidence': 'Synthetic full textless review only',
                                'targets': [p['image'] for p in selected]})
    enable_modules(root, route='B')
    return selected


class ExportTests(unittest.TestCase):
    def test_direct_export_functions_reject_unclassified_scope(self):
        for operation in (exports.export_slides, exports.export_documents, exports.collect):
            with self.subTest(operation=operation.__name__):
                root = fixture(register_scope=False)
                with self.assertRaisesRegex(ValueError, 'workflow|Workflow|scope|mode'):
                    operation(root)

    def test_selected_export_does_not_require_unused_empty_core_approval(self):
        root = fixture()
        unused = ['_state/' + name + '.json' for name in ('story', 'math', 'assets')]
        state.record_approval(root, {'decision': 'rejected', 'user_evidence': 'Unused empty records are not adopted',
            'targets': [{'path': path, 'sha256': state.sha256(root / path)} for path in unused]})
        try:
            result = exports.export_slides(root)
        except ValueError as exc:
            self.fail('Selected image export must only require its real input records: ' + str(exc))
        self.assertTrue((root / result['artifacts']['image-slides-pptx']['path']).exists())
        self.assertTrue(set(unused).isdisjoint(result['source_versions']))

    def test_selected_export_checks_used_math_reference(self):
        root = fixture()
        pages = state.read_json(root / '_state/pages.json')
        pages['pages'][0]['math_ids'] = ['MISSING']
        state.write_json(root / '_state/pages.json', pages)
        state.record_approval(root, {'user_evidence': 'Synthetic malformed content approval', 'targets': [
            {'path': '_state/pages.json', 'sha256': state.sha256(root / '_state/pages.json')}]})
        enable_modules(root)
        with self.assertRaisesRegex(ValueError, 'math|Math|MISSING'):
            exports.export_slides(root)
        self.assertFalse(list((root / 'slides').glob('export-v*')))

    def test_handoff_route_must_match_explicit_choice(self):
        root = fixture()
        selected = textless_selection(root)
        enable_modules(root, route='A')
        with self.assertRaisesRegex(ValueError, 'route|Route'):
            exports.export_handoff(root, {'route': 'B', 'pages': selected})
        self.assertFalse(list((root / 'editable/handoff').glob('handoff-v*')))

    def test_scoped_collect_excludes_unrequested_stale_outputs_and_assets(self):
        for module, prefix in [('documents', 'documents/export-v'), ('editable', 'editable/handoff/')]:
            with self.subTest(module=module):
                root = fixture()
                exports.export_slides(root)
                exports.export_documents(root)
                pages = sorted(state.read_json(root / '_state/pages.json')['pages'], key=lambda page: page['order'])
                exports.export_handoff(root, {'route': 'A', 'pages': pages})
                record = state.load_project(root)
                record['stale_targets'].append('image-slides-pdf')
                state.write_json(root / '_state/project.json', record)
                state.write_json(root / '_state/assets.json', {'assets': [{
                    'asset_id': 'UNRELATED', 'version': 'v001',
                    'files': [{'path': 'assets/not-part-of-this-task.png', 'sha256': '0' * 64}]}]})
                (root / 'slides/image-prompts.md').write_text('Unrelated historic prompt.', encoding='utf-8')
                enable_modules(root, modules=[module])
                try:
                    result = exports.collect(root)
                except ValueError as exc:
                    self.fail('Unrequested historic products must not block a scoped collection: ' + str(exc))
                self.assertTrue(result['artifacts'])
                self.assertTrue(all(item['path'].startswith(prefix) for item in result['artifacts'].values()))
                self.assertNotIn('_state/assets.json', result['source_versions'])

    def test_collect_requires_adoption_for_shared_assets_it_copies(self):
        root = fixture()
        path = root / 'assets/unapproved.png'
        Image.new('RGB', (32, 18), 'red').save(path)
        state.write_json(root / '_state/assets.json', {'assets': [{
            'asset_id': 'UNAPPROVED', 'version': 'v001',
            'files': [{'path': 'assets/unapproved.png', 'sha256': state.sha256(path)}]}]})
        with self.assertRaisesRegex(ValueError, 'confirmation'):
            exports.collect(root)
        self.assertFalse(list((root / 'deliveries').glob('delivery-v*')))

    def test_selected_documents_can_use_page_text_without_page_images(self):
        root = fixture()
        pages = state.read_json(root / '_state/pages.json')
        for page in pages['pages']:
            page.pop('image')
        state.write_json(root / '_state/pages.json', pages)
        source = {'_state/pages.json': state.sha256(root / '_state/pages.json')}
        state.record_approval(root, {'user_evidence': 'Synthetic source text adopted without page-image work',
            'targets': [{'path': path, 'sha256': digest} for path, digest in source.items()]})
        for kind in exports.DOCUMENTS:
            path = root / f'_state/documents/{kind}.json'
            record = state.read_json(path)
            record['source_versions'] = source
            state.write_json(path, record)
        enable_modules(root, modules=['documents'])
        try:
            result = exports.export_documents(root)
        except ValueError as exc:
            self.fail('Standalone documents do not require unrelated page-image production: ' + str(exc))
        self.assertEqual(set(result['documents']), set(exports.DOCUMENTS))

    def test_images_preserve_order_aspect_and_versions(self):
        root = fixture()
        report = exports.export_slides(root)
        self.assertEqual(report['page_order'], ['page-1', 'page-2'])
        deck = Presentation(root / report['artifacts']['image-slides-pptx']['path'])
        self.assertEqual([s.shapes[0].name for s in deck.slides], ['page-1', 'page-2'])
        for number, slide in enumerate(deck.slides, 1):
            self.assertEqual(hashlib.sha256(slide.shapes[0].image.blob).hexdigest(),
                             state.sha256(root / f'slides/fixture-{number}.png'))
        for shape, ratio in [(deck.slides[0].shapes[0], 16 / 9), (deck.slides[1].shapes[0], 4 / 3)]:
            self.assertAlmostEqual(shape.width / shape.height, ratio, places=5)
        second = report['image_layout'][1]['content_rect_pt']
        self.assertEqual(second, [120.0, 0.0, 720.0, 540.0])
        pdf = PdfReader(root / report['artifacts']['image-slides-pdf']['path'])
        self.assertEqual(len(pdf.pages), 2)
        self.assertEqual(list(pdf.pages[0].mediabox), [0, 0, 960, 540])
        first_hash = report['artifacts']['image-slides-pptx']['sha256']
        again = exports.export_slides(root)
        self.assertNotEqual(report['manifest'], again['manifest'])
        self.assertEqual(state.sha256(root / report['artifacts']['image-slides-pptx']['path']), first_hash)

    def test_three_documents_selectable_chinese_and_long_table(self):
        root = fixture()
        report = exports.export_documents(root)
        normalize = lambda text: ''.join(text.split())
        for kind, expected in [('classroom-script', '同学们，请观察第1组图形，比较大小。'),
                               ('lesson-presentation', '通过观察与表达建立数学关系。'),
                               ('lesson-plan', '第089步：引导学生比较并解释，保留符号 3 < 5 & 8 > 6。')]:
            artifacts = report['documents'][kind]['artifacts']
            docx = root / artifacts[kind + '-docx']['path']
            pdf = root / artifacts[kind + '-pdf']['path']
            markdown = root / artifacts[kind + '-md']['path']
            markdown_text = html.unescape(markdown.read_text(encoding='utf-8').replace('<br/>', '\n'))
            self.assertIn(normalize(expected), normalize(markdown_text))
            document = Document(docx)
            text = '\n'.join(p.text for p in document.paragraphs)
            text += '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
            self.assertIn(normalize(expected), normalize(text))
            reader = PdfReader(pdf)
            pdf_text = '\n'.join(p.extract_text() for p in reader.pages)
            self.assertIn(normalize(expected), normalize(pdf_text))
            self.assertAlmostEqual(float(reader.pages[0].mediabox.width), 595.2756, places=3)
            with zipfile.ZipFile(docx) as archive:
                styles = archive.read('word/styles.xml').decode('utf-8')
                self.assertIn('w:eastAsia="' + report['documents'][kind]['font_family'] + '"', styles)
                self.assertIn('w:eastAsia="zh-CN"', styles)
            if kind == 'classroom-script':
                self.assertEqual(document.paragraphs[1].text, '第01页  第1个数学任务')
                self.assertAlmostEqual(document.styles['Normal'].font.size.pt, 11)
                self.assertEqual(document.styles['Normal'].paragraph_format.line_spacing, 1.4)
            if kind == 'lesson-plan':
                self.assertGreater(len(reader.pages), 2)
                self.assertIn('第000步', normalize(reader.pages[0].extract_text()))
                self.assertEqual(len(document.tables[0].columns), 5)
                self.assertIn('tblHeader', document.tables[0].rows[0]._tr.xml)
                self.assertNotIn('cantSplit', document.tables[0]._tbl.xml)
                for page in reader.pages[1:]:
                    self.assertIn('教师活动', normalize(page.extract_text()))

    def test_rejects_changed_basis_unapproved_image_and_wrong_page_coverage(self):
        root = fixture()
        script = root / '_state/documents/classroom-script.json'
        data = state.read_json(script)
        data['content'][1]['page_id'] = 'invented-page'
        state.write_json(script, data)
        with self.assertRaisesRegex(ValueError, 'cover every page'):
            exports.export_documents(root)
        self.assertFalse(list((root / 'documents').glob('export-v*')))
        root = fixture()
        image = root / 'slides/fixture-1.png'
        Image.new('RGB', (800, 450), 'green').save(image)
        with self.assertRaisesRegex(ValueError, 'confirmation'):
            exports.export_slides(root)
        root = fixture()
        record = state.read_json(root / '_state/documents/lesson-plan.json')
        record['source_versions']['_state/math.json'] = '0' * 64
        state.write_json(root / '_state/documents/lesson-plan.json', record)
        with self.assertRaisesRegex(ValueError, 'Source version changed'):
            exports.export_documents(root)
        root = fixture()
        path = root / '_state/documents/lesson-presentation.json'
        record = state.read_json(path)
        record['content'][0]['page_ids'].append('invented-page')
        state.write_json(path, record)
        with self.assertRaisesRegex(ValueError, 'page references'):
            exports.export_documents(root)

    def test_collection_only_current_registered_files_and_stale_rejection(self):
        root = fixture()
        exports.export_slides(root)
        (root / 'inputs/private-secret.txt').write_text('fixture-private-marker')
        report = exports.collect(root)
        self.assertEqual(len(report['artifacts']), 4)
        delivery = (root / report['manifest']).parent
        self.assertFalse(any('private' in p.name for p in delivery.rglob('*')))
        project = state.load_project(root)
        project['stale_targets'].append('image-slides-pdf')
        state.write_json(root / '_state/project.json', project)
        with self.assertRaisesRegex(ValueError, 'Stale artifact'):
            exports.collect(root)
        root = fixture()
        report = exports.export_slides(root)
        source = root / '_state/pages.json'
        source.write_text(source.read_text(encoding='utf-8') + ' ', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'Source version changed|confirmation|Workflow blocked.*stale'):
            exports.collect(root)
        root = fixture()
        report = exports.export_slides(root)
        artifact = root / report['artifacts']['image-slides-pdf']['path']
        with artifact.open('ab') as stream:
            stream.write(b'changed fixture')
        with self.assertRaisesRegex(ValueError, 'Artifact hash changed'):
            exports.collect(root)

    def test_markdown_literal_table_cells_roundtrip(self):
        text = '比较 | 数量\n保留 <br/> 与 & 和 \\ * 符号'
        rendered = exports._markdown([('title', '测试'), ('table', [list(exports.COLUMNS), [text] * 5])])
        rows = [line for line in rendered.splitlines() if line.startswith('|')]
        self.assertEqual(len(rows), 3)
        cells = rows[-1].split('|')[1:-1]
        self.assertEqual(len(cells), 5)
        for cell in cells:
            self.assertEqual(html.unescape(cell.strip().replace('<br/>', '\n')), text)
        self.assertNotIn('\n\n| ---', rendered)

    def test_handoff_b_preserves_canonical_and_selected_image_identity(self):
        root = fixture()
        original = state.sha256(root / '_state/pages.json')
        selected = textless_selection(root)
        report = exports.export_handoff(root, {'route': 'B', 'pages': selected})
        self.assertEqual(state.sha256(root / '_state/pages.json'), original)
        self.assertEqual(report['page_order'], ['page-1', 'page-2'])
        self.assertEqual(report['review_status'], 'waiting_manual_canva')
        deck = Presentation(root / report['artifacts']['canva-handoff-b-pptx']['path'])
        pdf = PdfReader(root / report['artifacts']['canva-handoff-b-pdf']['path'])
        self.assertEqual(len(deck.slides), len(selected))
        self.assertEqual(len(pdf.pages), len(selected))
        for number, (slide, item, mapping) in enumerate(zip(deck.slides, selected, report['page_mapping']), 1):
            self.assertEqual(len(slide.shapes), 1)
            self.assertEqual(slide.shapes[0].name, item['page_id'])
            self.assertEqual(hashlib.sha256(slide.shapes[0].image.blob).hexdigest(), item['image']['sha256'])
            self.assertEqual(state.sha256(root / mapping['image_copy']), item['image']['sha256'])
            self.assertEqual(mapping['page_id'], item['page_id'])
            self.assertEqual(mapping['handoff_slide_number'], number)
            with Image.open(root / item['image']['path']) as source:
                self.assertAlmostEqual(slide.shapes[0].width / slide.shapes[0].height,
                                       source.width / source.height, places=5)
                self.assertEqual(list(pdf.pages[number - 1].images)[0].image.convert('RGB').tobytes(),
                                 source.convert('RGB').tobytes())
        collected = exports.collect(root)
        self.assertTrue(any(a['path'] == report['manifest'] for a in collected['artifacts'].values()))
        self.assertTrue(any(a['path'] == report['page_mapping'][0]['image_copy'] for a in collected['artifacts'].values()))
        pages = sorted(state.read_json(root / '_state/pages.json')['pages'], key=lambda p: p['order'])
        with self.assertRaisesRegex(ValueError, 'separate confirmed'):
            exports.export_handoff(root, {'route': 'B', 'pages': pages})
        enable_modules(root, route='A')
        with self.assertRaisesRegex(ValueError, 'ordered subset'):
            exports.export_handoff(root, {'route': 'A', 'pages': list(reversed(pages))})
        report_a = exports.export_handoff(root, {'route': 'A', 'pages': [pages[0]]})
        self.assertEqual(report_a['route'], 'A')
        with (root / selected[1]['image']['path']).open('ab') as stream:
            stream.write(b'changed')
        with self.assertRaisesRegex(ValueError, 'Source version changed'):
            exports.collect(root)
        root = fixture()
        page_record = state.read_json(root / '_state/pages.json')
        pending_page = next(p for p in page_record['pages'] if p['page_id'] == 'page-1')
        pending_page['image'] = None
        state.write_json(root / '_state/pages.json', page_record)
        state.record_approval(root, {'user_evidence': 'Synthetic partial sample approval',
                                    'targets': [{'path': '_state/pages.json', 'sha256': state.sha256(root / '_state/pages.json')}]})
        enable_modules(root, route='A')
        sample = next(p for p in page_record['pages'] if p['page_id'] == 'page-2')
        self.assertEqual(exports.export_handoff(root, {'route': 'A', 'pages': [sample]})['page_order'], ['page-2'])

    def test_handoff_b_rejects_missing_pages_before_creating_outputs(self):
        root = fixture()
        selected = textless_selection(root)
        with self.assertRaisesRegex(ValueError, 'every canonical page'):
            exports.export_handoff(root, {'route': 'B', 'pages': selected[1:]})
        self.assertFalse(list((root / 'editable/handoff').glob('handoff-v*')))

    def test_handoff_b_rejects_duplicate_unknown_and_misordered_pages(self):
        root = fixture()
        selected = textless_selection(root)
        for items in ([selected[0], selected[0]], list(reversed(selected)),
                      [selected[0], {'page_id': 'unknown', 'image': selected[1]['image']}]):
            with self.subTest(page_ids=[p['page_id'] for p in items]):
                with self.assertRaisesRegex(ValueError, 'ordered subset'):
                    exports.export_handoff(root, {'route': 'B', 'pages': items})
        self.assertFalse(list((root / 'editable/handoff').glob('handoff-v*')))

    def test_collect_registered_cover_reference_prompt_without_changing_sources(self):
        root = fixture()
        cover_dir = root / 'slides/covers/example/v001'
        cover_dir.mkdir(parents=True)
        image = cover_dir / 'cover.png'
        prompt = cover_dir / 'prompt.txt'
        Image.new('RGB', (80, 45), 'green').save(image)
        prompt.write_text('Synthetic adopted cover prompt.', encoding='utf-8')
        asset = {'asset_id': 'REF001', 'kind': 'style_reference',
                 'prompt_path': prompt.relative_to(root).as_posix(),
                 'files': [{'path': image.relative_to(root).as_posix(), 'sha256': state.sha256(image)}]}
        state.write_json(root / '_state/assets.json', {'assets': [asset]})
        state.record_approval(root, {'user_evidence': 'Synthetic cover fixture only', 'targets': [
            {'path': '_state/assets.json', 'sha256': state.sha256(root / '_state/assets.json')}]})
        basis = exports._basis(root)[1]
        prompt_basis = {**basis, '_state/assets.json': state.sha256(root / '_state/assets.json')}
        state.write_json(root / '_state/prompt-exports.json', {'source_versions': prompt_basis, 'files': [
            {'path': asset['prompt_path'], 'sha256': state.sha256(prompt)}]})
        report = exports.collect(root)
        copies = [a for a in report['artifacts'].values() if a['path'] == asset['prompt_path']]
        self.assertEqual(len(copies), 1)
        self.assertEqual(state.sha256(root / copies[0]['delivery_path']), state.sha256(prompt))
        self.assertEqual(exports._basis(root)[1], basis)
        for kind, relative in [('character', asset['prompt_path']),
                               ('style_reference', 'slides/covers/unrelated/prompt.txt')]:
            with self.subTest(kind=kind, path=relative):
                changed = {**asset, 'kind': kind, 'prompt_path': relative}
                other_prompt = root / relative
                if not other_prompt.exists():
                    other_prompt.parent.mkdir(parents=True, exist_ok=True)
                    other_prompt.write_text('Unrelated synthetic prompt.', encoding='utf-8')
                state.write_json(root / '_state/assets.json', {'assets': [changed]})
                state.record_approval(root, {'user_evidence': 'Synthetic rejection fixture only', 'targets': [
                    {'path': '_state/assets.json', 'sha256': state.sha256(root / '_state/assets.json')}]})
                with self.assertRaisesRegex(ValueError, 'Asset prompt_path'):
                    exports.collect(root)
        self.assertEqual(len(list((root / 'deliveries').glob('delivery-v*'))), 1)

    def test_collect_shared_assets_and_current_prompt_manifest_only(self):
        root = fixture()
        asset_image = root / 'assets/characters/reference.png'
        Image.new('RGB', (100, 100), 'yellow').save(asset_image)
        prompt = root / 'assets/characters/reference-prompt.md'
        prompt.write_text('合成测试共享角色提示词。', encoding='utf-8')
        asset_record = {'assets': [{'asset_id': 'hero', 'version': 'v001',
                        'prompt_path': 'assets/characters/reference-prompt.md',
                        'files': [{'path': 'assets/characters/reference.png', 'sha256': state.sha256(asset_image)}]}]}
        state.write_json(root / '_state/assets.json', asset_record)
        state.record_approval(root, {'user_evidence': 'Synthetic asset approval',
                                    'targets': [{'path': '_state/assets.json', 'sha256': state.sha256(root / '_state/assets.json')}]})
        paths = ['slides/image-prompts.md', 'planning/visible-text.md', 'assets/asset-prompts.md']
        for path in paths:
            (root / path).write_text('合成测试提示词文案', encoding='utf-8')
        paths.append('assets/characters/reference-prompt.md')
        manifest = {'source_versions': {p: state.sha256(root / p) for p in exports.CORE},
                    'files': [{'path': p, 'sha256': state.sha256(root / p)} for p in paths]}
        state.write_json(root / '_state/prompt-exports.json', manifest)
        (root / 'assets/characters/unregistered-secret.txt').write_text('synthetic private text')
        exported = exports.collect(root)
        delivered = [a['path'] for a in exported['artifacts'].values()]
        for path in [*paths, 'assets/characters/reference.png']:
            self.assertIn(path, delivered)
        self.assertNotIn('assets/characters/unregistered-secret.txt', delivered)
        prompt.write_text('changed source', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'Artifact hash changed'):
            exports.collect(root)
        root = fixture()
        (root / 'slides/image-prompts.md').write_text('untracked prompt', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'generation manifest'):
            exports.collect(root)


if __name__ == '__main__':
    unittest.main()
