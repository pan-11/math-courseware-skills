"""Deterministic, versioned exports of approved images and canonical teaching text."""
from pathlib import Path
from collections import Counter
import os
import shutil
from xml.sax.saxutils import escape

from . import state

DOCUMENTS = ('classroom-script', 'lesson-presentation', 'lesson-plan')
CORE = ('_state/pages.json', '_state/story.json', '_state/math.json', '_state/assets.json')
COLUMNS = ('环节', '教师活动', '学生活动', '任务与评价', '设计意图/二次备课')


def _version_dir(project, parent, prefix='export-v'):
    base = state.resolve(project, parent)
    base.mkdir(parents=True, exist_ok=True)
    number = 1
    while (base / f'{prefix}{number:03d}').exists():
        number += 1
    output = base / f'{prefix}{number:03d}'
    output.mkdir()
    return output


def _relative(project, path):
    return Path(path).relative_to(Path(project).resolve()).as_posix()


def _check_versions(project, versions):
    if not isinstance(versions, dict) or not versions:
        raise ValueError('Nonempty source_versions required')
    for path, expected in versions.items():
        source = state.resolve(project, path)
        if not source.is_file() or state.sha256(source) != expected:
            raise ValueError('Source version changed: ' + path)


def _basis(project, image_page_ids=None):
    state.load_project(project)
    state.require_approved(project, *CORE)
    pages = state.read_json(state.resolve(project, CORE[0]))['pages']
    if not pages:
        raise ValueError('Approved pages must not be empty')
    ids = [p['page_id'] for p in pages]
    orders = [p['order'] for p in pages]
    if len(set(ids)) != len(ids) or any(not isinstance(x, str) or not x for x in ids):
        raise ValueError('Unique stable page IDs required')
    if any(type(x) is not int or x < 1 for x in orders) or len(set(orders)) != len(orders):
        raise ValueError('Unique positive page order required')
    pages = sorted(pages, key=lambda p: p['order'])
    versions = {path: state.sha256(state.resolve(project, path)) for path in CORE}
    for page in pages:
        if image_page_ids is not None and page['page_id'] not in image_page_ids:
            continue
        image = page.get('image', {})
        if not image.get('path') or not image.get('sha256'):
            raise ValueError('Approved image required for ' + page['page_id'])
        state.require_approved(project, image['path'])
        _check_versions(project, {image['path']: image['sha256']})
        versions[image['path']] = image['sha256']
    return pages, versions


def _register(project, output, artifact_id, paths, versions, page_ids, extra=None):
    _check_versions(project, versions)
    artifacts = {}
    for path in paths:
        key = artifact_id + '-' + path.suffix.lstrip('.')
        relative = _relative(project, path)
        metadata = {'source_versions': versions, 'page_order': page_ids,
                    'export_kind': artifact_id, **(extra or {})}
        artifacts[key] = state.register_artifact(
            project, key, relative, [*versions, *page_ids], metadata)
    report = {'schema_version': state.SCHEMA, 'created_at': state.now(),
              'artifacts': artifacts, 'source_versions': versions, 'page_order': page_ids,
              'visual_review': 'WPS display not verified', **(extra or {})}
    manifest = output / 'manifest.json'
    state.write_json(manifest, report)
    report['manifest'] = _relative(project, manifest)
    return report


def _render_images(project, pages, output, name):
    """Shared image placement for the full deck and explicit A/B handoff selections."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Pt
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader

    dimensions = []
    for page in pages:
        with Image.open(state.resolve(project, page['image']['path'])) as image:
            image.verify()
        with Image.open(state.resolve(project, page['image']['path'])) as image:
            dimensions.append(image.size)
    pptx, pdf = output / (name + '.pptx'), output / (name + '.pdf')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(960), Pt(540)
    canvas = Canvas(str(pdf), pagesize=(960, 540))
    rects = []
    for page, (width, height) in zip(pages, dimensions):
        scale = min(960 / width, 540 / height)
        w, h = width * scale, height * scale
        x, y = (960 - w) / 2, (540 - h) / 2
        source = state.resolve(project, page['image']['path'])
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_picture(str(source), Pt(x), Pt(y), Pt(w), Pt(h))
        shape.name = page['page_id']
        canvas.drawImage(str(source), x, 540 - y - h, w, h, mask='auto')
        canvas.showPage()
        rects.append({'page_id': page['page_id'], 'source_size_px': [width, height],
                      'content_rect_pt': [x, y, w, h], 'fit': 'letterbox'})
    deck.save(pptx)
    canvas.save()
    if len(Presentation(pptx).slides) != len(pages) or len(PdfReader(pdf).pages) != len(pages):
        raise ValueError('Export page count mismatch')
    return [pptx, pdf], rects


def export_slides(project):
    """Export existing approved images as a 960x540 pt image deck and matching PDF."""
    project = Path(project).resolve()
    pages, versions = _basis(project)
    output = _version_dir(project, 'slides')
    paths, rects = _render_images(project, pages, output, 'lesson-images')
    return _register(project, output, 'image-slides', paths, versions,
                     [p['page_id'] for p in pages], {'image_layout': rects})


def export_handoff(project, selection):
    """Export authorized A pages or all ordered B pages without rewriting canonical images."""
    project = Path(project).resolve()
    route, chosen = selection.get('route'), selection.get('pages')
    if route not in ('A', 'B') or not isinstance(chosen, list) or not chosen:
        raise ValueError('Handoff requires route A/B and nonempty pages selection')
    ids = [item['page_id'] for item in chosen]
    pages, versions = _basis(project, ids)
    canonical = {p['page_id']: p for p in pages}
    if len(set(ids)) != len(ids) or ids != [p['page_id'] for p in pages if p['page_id'] in ids]:
        raise ValueError('Handoff pages must be an ordered subset of canonical page IDs')
    if route == 'B' and ids != [p['page_id'] for p in pages]:
        raise ValueError('Route B requires every canonical page exactly once in page order')
    for item in chosen:
        image = item['image']
        state.require_approved(project, image['path'])
        _check_versions(project, {image['path']: image['sha256']})
        original = canonical[item['page_id']]['image']
        if route == 'A' and any(image[key] != original[key] for key in ('path', 'sha256')):
            raise ValueError('Route A uses the canonical confirmed text-bearing image')
        if route == 'B' and state.resolve(project, image['path']) == state.resolve(project, original['path']):
            raise ValueError('Route B must use a separate confirmed textless image')
        versions[image['path']] = image['sha256']
    output = _version_dir(project, 'editable/handoff', 'handoff-v')
    paths, rects = _render_images(project, chosen, output, 'canva-handoff')
    mapping, files = [], []
    for number, item in enumerate(chosen, 1):
        source = state.resolve(project, item['image']['path'])
        copy = output / (f'page-{number:03d}' + source.suffix.lower())
        shutil.copy2(source, copy)
        if state.sha256(copy) != item['image']['sha256']:
            raise ValueError('Handoff image copy hash mismatch')
        paths.append(copy)
        mapping.append({'page_id': item['page_id'], 'handoff_slide_number': number,
                        'canonical_image': canonical[item['page_id']]['image'],
                        'selected_image': item['image'], 'image_copy': _relative(project, copy)})
    files = [{'path': _relative(project, path), 'sha256': state.sha256(path)} for path in paths]
    artifact_id = 'canva-handoff-' + route.lower()
    report = _register(project, output, artifact_id, paths[:2], versions, ids,
                       {'route': route, 'review_status': 'waiting_manual_canva',
                        'page_mapping': mapping, 'image_layout': rects, 'files': files})
    manifest_key = artifact_id + '-manifest'
    report['artifacts'][manifest_key] = state.register_artifact(
        project, manifest_key, report['manifest'], [*versions, *ids],
        {'source_versions': versions, 'handoff_manifest': True, 'page_order': ids})
    return report


def _text(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError('Nonempty text required: ' + label)
    if any(ord(c) < 32 and c not in '\n\t\r' for c in value):
        raise ValueError('Invalid control character: ' + label)
    return value


def _sections(sections):
    if not isinstance(sections, list) or not sections:
        raise ValueError('Nonempty content sections required')
    for item in sections:
        _text(item.get('heading'), 'heading')
        _paragraphs(item.get('paragraphs'))


def _paragraphs(paragraphs):
    if not isinstance(paragraphs, list) or not paragraphs:
        raise ValueError('Nonempty paragraphs required')
    for text in paragraphs:
        _text(text, 'paragraph')


def _document(project, kind, pages, versions):
    relative = f'_state/documents/{kind}.json'
    source = state.resolve(project, relative)
    record = state.read_json(source)
    _text(record.get('title'), 'title')
    locked = record.get('source_versions')
    _check_versions(project, locked)
    if any(locked.get(path) != digest for path, digest in versions.items()):
        raise ValueError('Document must lock approved story/math/assets/pages and page images')
    ids = [p['page_id'] for p in pages]
    if record.get('page_order') != ids:
        raise ValueError('Document page_order differs from current approved page order')
    content = record.get('content')
    refs = []
    if kind == 'classroom-script':
        if not isinstance(content, list) or [x.get('page_id') for x in content] != ids:
            raise ValueError('Classroom script must cover every page once in approved order')
        for item in content:
            _text(item.get('title'), 'page title')
            _paragraphs(item.get('paragraphs'))
    else:
        if kind == 'lesson-presentation':
            _sections(content)
            entries = content
        else:
            if not isinstance(content, dict):
                raise ValueError('Lesson plan content must be an object')
            _sections(content.get('metadata'))
            entries = content.get('rows')
            if not isinstance(entries, list) or not entries:
                raise ValueError('Lesson plan requires process rows')
            for row in entries:
                if not isinstance(row.get('cells'), list) or len(row['cells']) != 5:
                    raise ValueError('Lesson plan rows need five cells')
                for cell in row['cells']:
                    _text(cell, 'table cell')
            if content.get('afterword'):
                _sections(content['afterword'])
        for item in entries:
            if not isinstance(item.get('page_ids'), list):
                raise ValueError('Internal page_ids association required')
            refs.extend(item['page_ids'])
        if set(refs) != set(ids):
            raise ValueError('Document page references contain missing or unknown page IDs')
    return record, {**locked, relative: state.sha256(source)}


def _font():
    """Use installed fonts only; never copy font files into release packages."""
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase import pdfmetrics
    fonts = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
    candidates = [(fonts / 'STSONG.TTF', '华文宋体'), (fonts / 'simfang.ttf', '仿宋'),
                  (fonts / 'simkai.ttf', '楷体'),
                  (Path('/usr/share/fonts/truetype/arphic/uming.ttc'), 'AR PL UMing CN')]
    for path, family in candidates:
        if path.is_file():
            font = TTFont('CoursewareCJK', str(path))
            if all(ord(char) in font.face.charToGlyph for char in '教师数学课堂活动'):
                pdfmetrics.registerFont(font)
                return family, path, font
    raise ValueError('No supported local CJK TrueType font; provide an installed CJK font')


def _blocks(kind, record):
    yield ('title', record['title'])
    if kind == 'classroom-script':
        for number, item in enumerate(record['content'], 1):
            yield ('heading', f'第{number:02d}页  {item["title"]}')
            for paragraph in item['paragraphs']:
                yield ('paragraph', paragraph)
    else:
        sections = record['content'] if kind == 'lesson-presentation' else record['content']['metadata']
        for item in sections:
            yield ('heading', item['heading'])
            for paragraph in item['paragraphs']:
                yield ('paragraph', paragraph)
        if kind == 'lesson-plan':
            yield ('heading', '教学过程')
            yield ('table', [list(COLUMNS)] + [row['cells'] for row in record['content']['rows']])
            for item in record['content'].get('afterword', []):
                yield ('heading', item['heading'])
                for paragraph in item['paragraphs']:
                    yield ('paragraph', paragraph)


def _render_document(output, kind, record, family, font):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle
    from pypdf import PdfReader

    blocks = list(_blocks(kind, record))
    strings = [str(cell) for block, value in blocks for row in
               (value if block == 'table' else [[value]]) for cell in row]
    missing = sorted({c for text in strings for c in text if not c.isspace()
                      and ord(c) not in font.face.charToGlyph})
    if missing:
        raise ValueError('Installed PDF font lacks characters: ' + ''.join(missing))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(23)
    section.top_margin = section.bottom_margin = Mm(22)
    for name, size in [('Normal', 11), ('Title', 18), ('Heading 1', 13)]:
        style = doc.styles[name]
        style.font.name, style.font.size = family, Pt(size)
        rpr = style.element.get_or_add_rPr()
        rpr.rFonts.set(qn('w:eastAsia'), family)
        language = OxmlElement('w:lang')
        language.set(qn('w:eastAsia'), 'zh-CN')
        rpr.append(language)
        style.paragraph_format.line_spacing = 1.4
        style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = record['title']
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    page_field = OxmlElement('w:fldSimple')
    page_field.set(qn('w:instr'), 'PAGE')
    footer._p.append(page_field)
    styles = {key: ParagraphStyle(key, fontName='CoursewareCJK', fontSize=size,
               leading=leading, wordWrap='CJK', spaceAfter=5,
               keepWithNext=key in ('title', 'heading'))
              for key, size, leading in [('title', 18, 25), ('heading', 13, 18),
                                        ('paragraph', 11, 15.4), ('cell', 10.5, 13)]}
    def paragraph(text, style):
        return Paragraph(escape(text).replace('\n', '<br/>'), styles[style])
    flow = []
    widths_mm = [18, 42, 34, 35, 35]
    for block, value in blocks:
        if block != 'table':
            doc.add_paragraph(value, style={'title': 'Title', 'heading': 'Heading 1',
                                          'paragraph': 'Normal'}[block])
            item = paragraph(value, block)
            # KeepTogether around a heading plus a multi-page LongTable pushes the
            # whole process section to a fresh page. Let the table split in place.
            if block == 'heading' and kind == 'lesson-plan' and value == '教学过程':
                item.keepWithNext = False
            flow.append(item)
            continue
        table = doc.add_table(rows=0, cols=5)
        table.style, table.autofit = 'Table Grid', False
        table_width = table._tbl.tblPr.find(qn('w:tblW'))
        table_width.set(qn('w:type'), 'dxa')
        table_width.set(qn('w:w'), str(round(sum(widths_mm) / 25.4 * 1440)))
        margins = OxmlElement('w:tblCellMar')
        for side in ('top', 'left', 'bottom', 'right'):
            margin = OxmlElement('w:' + side)
            margin.set(qn('w:w'), '60')
            margin.set(qn('w:type'), 'dxa')
            margins.append(margin)
        table._tbl.tblPr.append(margins)
        for column, width in zip(table.columns, widths_mm):
            column.width = Mm(width)
        for index, values in enumerate(value):
            row = table.add_row()
            if index == 0:
                repeat = OxmlElement('w:tblHeader')
                row._tr.get_or_add_trPr().append(repeat)
            for cell, text, width in zip(row.cells, values, widths_mm):
                cell.width, cell.text = Mm(width), text
                if index == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'EDF2F7')
                    cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.2
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(10.5)
        flow.append(LongTable([[paragraph(cell, 'cell') for cell in row] for row in value],
                    colWidths=[w * mm for w in widths_mm], repeatRows=1,
                    splitByRow=0, splitInRow=1, hAlign='LEFT',
                    style=TableStyle([('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#ADB5BD')),
                                      ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EDF2F7')),
                                      ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                      ('LEFTPADDING', (0, 0), (-1, -1), 3),
                                      ('RIGHTPADDING', (0, 0), (-1, -1), 3)])))
    docx_path, pdf_path = output / (kind + '.docx'), output / (kind + '.pdf')
    doc.save(docx_path)
    def page_marks(canvas, document):
        canvas.saveState()
        canvas.setFont('CoursewareCJK', 9)
        canvas.drawString(23 * mm, A4[1] - 14 * mm, record['title'])
        canvas.drawRightString(A4[0] - 23 * mm, 13 * mm, str(document.page))
        canvas.restoreState()
    SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=23 * mm,
                      rightMargin=23 * mm, topMargin=22 * mm, bottomMargin=22 * mm,
                      title=record['title']).build(flow, onFirstPage=page_marks, onLaterPages=page_marks)
    loaded = Document(docx_path)
    word_text = '\n'.join(p.text for p in loaded.paragraphs)
    word_text += '\n' + '\n'.join(c.text for t in loaded.tables for r in t.rows for c in r.cells)
    reader = PdfReader(pdf_path)
    pdf_text = '\n'.join(p.extract_text() or '' for p in reader.pages)
    normalize = lambda text: ''.join(text.split())
    # PDF reading order interleaves cells when a row splits across pages. For tables,
    # use character coverage; contiguous paragraph checks remain exact elsewhere.
    pdf_counts = Counter(normalize(pdf_text))
    expected_counts = Counter(normalize(''.join(strings)))
    if expected_counts - pdf_counts:
        raise ValueError('PDF content character coverage mismatch')
    for text in strings:
        if normalize(text) not in normalize(word_text):
            raise ValueError('DOCX content readback mismatch')
    for block, text in blocks:
        if block != 'table' and normalize(text) not in normalize(pdf_text):
            raise ValueError('PDF paragraph readback mismatch: ' + text[:30])
    markdown_path = output / (kind + '.md')
    markdown_path.write_text(_markdown(blocks), encoding='utf-8')
    return [docx_path, pdf_path, markdown_path], len(reader.pages)


def _markdown(blocks):
    # Encode Markdown punctuation as entities so literal pipes, backslashes,
    # HTML-looking text and line breaks survive both prose and GFM table cells.
    def literal(text):
        value = ''.join(f'&#{ord(char)};' if char in '\\`*_{}[]()#+-.!|~'
                        else escape(char) for char in text)
        return value.replace('\r\n', '\n').replace('\r', '\n').replace('\n', '<br/>')
    lines = []
    for kind, value in blocks:
        if kind == 'table':
            for number, row in enumerate(value):
                lines.append('| ' + ' | '.join(literal(cell) for cell in row) + ' |')
                if number == 0:
                    lines.append('| ' + ' | '.join('---' for _ in row) + ' |')
        else:
            lines.append({'title': '# ', 'heading': '## ', 'paragraph': ''}[kind] + literal(value))
        lines.append('')
    return '\n'.join(lines)


def export_documents(project):
    """Render all three validated canonical records; final user approval is not required."""
    project = Path(project).resolve()
    pages, versions = _basis(project)
    records = {kind: _document(project, kind, pages, versions) for kind in DOCUMENTS}
    family, font_path, font = _font()
    output = _version_dir(project, 'documents')
    reports = {}
    for kind, (record, sources) in records.items():
        folder = output / kind
        folder.mkdir()
        paths, count = _render_document(folder, kind, record, family, font)
        reports[kind] = _register(project, folder, kind, paths, sources,
                                 [p['page_id'] for p in pages],
                                 {'pdf_pages': count, 'font_family': family,
                                  'font_file_name': font_path.name,
                                  'canonical_sha256': sources[f'_state/documents/{kind}.json']})
    result = {'documents': reports, 'output': _relative(project, output)}
    state.write_json(output / 'manifest.json', result)
    return result


def collect(project):
    """Collect explicit current artifacts, canonical media and verified prompt manifests."""
    project = Path(project).resolve()
    pages, basis = _basis(project)
    record = state.load_project(project)
    artifacts = record.get('artifacts', {})
    selected = {}
    def add(key, artifact):
        path = state.resolve(project, artifact['path'])
        if key in record.get('stale_targets', []):
            raise ValueError('Stale artifact requires regeneration: ' + key)
        if not path.is_file() or state.sha256(path) != artifact['sha256']:
            raise ValueError('Artifact hash changed: ' + key)
        if not artifact.get('source_versions'):
            raise ValueError('Artifact has no verifiable source_versions: ' + key)
        _check_versions(project, artifact['source_versions'])
        selected[key] = artifact

    for key, artifact in artifacts.items():
        path = state.resolve(project, artifact['path'])
        relative = _relative(project, path)
        if artifact.get('handoff_manifest'):
            if not relative.startswith('editable/handoff/handoff-v') or path.name != 'manifest.json':
                raise ValueError('Handoff manifest must be inside a versioned handoff folder')
            add(key, artifact)
            manifest = state.read_json(path)
            if manifest.get('source_versions') != artifact['source_versions']:
                raise ValueError('Handoff manifest source map differs from registered metadata')
            if not manifest.get('files'):
                raise ValueError('Handoff manifest has no files')
            for number, item in enumerate(manifest['files']):
                file = state.resolve(project, item['path'])
                if file.parent != path.parent or file.suffix.lower() not in ('.pptx', '.pdf', '.png', '.jpg', '.jpeg', '.webp'):
                    raise ValueError('Handoff manifest file escapes its explicit package')
                add(f'{key}-file-{number}', {**item, 'source_versions': artifact['source_versions']})
            continue
        if path.suffix.lower() not in ('.pptx', '.pdf', '.docx', '.md'):
            continue
        if not (relative.startswith('slides/export-v') or relative.startswith('documents/export-v')
                or relative.startswith('editable/output/') or relative.startswith('editable/handoff/handoff-v')):
            continue
        add(key, artifact)

    from PIL import Image
    def media(key, item, destination):
        path = state.resolve(project, item['path'])
        with Image.open(path) as image:
            image.verify()
        add(key, {**item, 'source_versions': basis, 'destination': destination + path.suffix.lower()})
    for number, page in enumerate(pages, 1):
        media('page-image-' + page['page_id'], page['image'], f'slides/page-images/page-{number:03d}')
    assets = state.read_json(state.resolve(project, '_state/assets.json')).get('assets', [])
    prompt_paths = {'slides/image-prompts.md', 'planning/visible-text.md', 'assets/asset-prompts.md'}
    for index, asset in enumerate(assets, 1):
        for number, item in enumerate(asset.get('files', []), 1):
            media(f'shared-asset-{index}-{number}', item, f'assets/shared/asset-{index:03d}-file-{number:03d}')
        if asset.get('prompt_path'):
            path = state.resolve(project, asset['prompt_path'])
            relative = _relative(project, path)
            cover_reference = (asset.get('kind') == 'style_reference' and relative.startswith('slides/covers/')
                               and any(path.parent == state.resolve(project, item['path']).parent
                                       for item in asset.get('files', [])))
            if (not relative.startswith('assets/') and not cover_reference) or path.suffix.lower() not in ('.md', '.txt') or any(
                    part.startswith('.') for part in Path(relative).parts):
                raise ValueError('Asset prompt_path must name an explicit text file under assets/ or beside its registered cover reference')
            if not path.is_file():
                raise ValueError('Missing canonical asset prompt: ' + relative)
            prompt_paths.add(relative)
    available = sorted(p for p in prompt_paths if state.resolve(project, p).is_file())
    if available:
        manifest_path = state.resolve(project, '_state/prompt-exports.json')
        if not manifest_path.is_file():
            raise ValueError('Prompt files need a current generation manifest')
        manifest = state.read_json(manifest_path)
        sources = manifest.get('source_versions')
        _check_versions(project, sources)
        if any(sources.get(path) != basis[path] for path in CORE):
            raise ValueError('Prompt manifest must lock current core records')
        files = {item['path']: item for item in manifest.get('files', [])}
        for number, relative in enumerate(available):
            if relative not in files:
                raise ValueError('Prompt missing from generation manifest: ' + relative)
            add(f'prompt-{number}', {**files[relative], 'source_versions': sources})
    if not selected:
        raise ValueError('No current registered deliverables')
    output = _version_dir(project, 'deliveries', 'delivery-v')
    copied = {}
    for key, artifact in selected.items():
        destination = (output / artifact.get('destination', artifact['path'])).resolve()
        if not destination.is_relative_to(output):
            raise ValueError('Delivery path escapes versioned output')
        destination.parent.mkdir(parents=True, exist_ok=True)
        if not destination.exists():
            shutil.copy2(state.resolve(project, artifact['path']), destination)
        if state.sha256(destination) != artifact['sha256']:
            raise ValueError('Delivery copy hash mismatch')
        copied[key] = {**artifact, 'delivery_path': _relative(project, destination)}
    report = {'created_at': state.now(), 'artifacts': copied,
              'source_versions': basis, 'page_order': [page['page_id'] for page in pages],
              'visual_review': 'WPS display not verified'}
    _check_versions(project, basis)
    for artifact in selected.values():
        _check_versions(project, artifact['source_versions'])
    state.write_json(output / 'manifest.json', report)
    report['manifest'] = _relative(project, output / 'manifest.json')
    return report
