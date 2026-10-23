"""Read-only, source-bound teaching materials and explicit collection manifests."""
from pathlib import Path, PureWindowsPath
import re

from . import state

TEACHERS = ('classroom-script', 'lesson-presentation', 'lesson-plan')
ROLES = (*TEACHERS, 'worksheet', 'blackboard', 'interactive')
DELIVERY_SUFFIXES = {'.pptx', '.pdf', '.docx', '.md', '.txt', '.html', '.htm', '.js', '.css',
                     '.json', '.png', '.jpg', '.jpeg', '.webp', '.svg', '.gif', '.ico', '.csv',
                     '.woff', '.woff2', '.ttf', '.otf', '.mp4', '.m4v', '.mov', '.webm',
                     '.mp3', '.wav', '.ogg', '.m4a'}
PRIVATE_NAMES = {'key', 'secret', 'secrets', 'credential', 'credentials', 'token', 'password',
                 'apikey', 'privatekey', 'idrsa', 'system', 'system32'}


def _delivery_path(relative):
    path = PureWindowsPath(relative)
    if path.suffix.lower() not in DELIVERY_SUFFIXES or any(
            part.startswith('.') or re.sub(r'[-_.\s]+', '', PureWindowsPath(part).stem.casefold()) in PRIVATE_NAMES
            for part in path.parts):
        raise ValueError('materials: private, hidden or unsupported delivery path: ' + str(relative))


def _file(project, item, label):
    if not isinstance(item, dict) or not item.get('path') or not item.get('sha256'):
        raise ValueError(label + ': actual path and sha256 required')
    path = state.resolve(project, item['path'])
    if not path.is_file() or not path.stat().st_size or state.sha256(path) != item['sha256']:
        raise ValueError(label + ': missing, empty or changed file: ' + item['path'])
    return path


def _sources(project, versions, required, label):
    if not isinstance(versions, dict) or not versions:
        raise ValueError(label + ': nonempty source_versions required')
    for path, sha in versions.items():
        _file(project, {'path': path, 'sha256': sha}, label + '/source')
    normalized = {state.resolve(project, path): sha for path, sha in versions.items()}
    if any(normalized.get(state.resolve(project, path)) != sha for path, sha in required.items()):
        raise ValueError(label + ': current course/deck source binding required')


def current_inputs(project, data, module):
    files = data.get('inputs', {}).get(module, {}).get('files', {})
    if not files:
        raise ValueError('materials: current module input required: ' + module)
    for item in files.values():
        _file(project, item, 'materials/current-input')
    return {item['path']: item['sha256'] for item in files.values()}


def artifact_role(key, artifact):
    role = artifact.get('role', artifact.get('export_kind'))
    if role in ROLES:
        return role
    return key if key in ROLES else next((name for name in ROLES if key.startswith(name + '-')), None)


def _populated(path, role):
    suffix = path.suffix.lower()
    if role in (*TEACHERS, 'worksheet'):
        if suffix == '.docx':
            from docx import Document
            doc = Document(path)
            return bool(doc.tables or any(p.text.strip() for p in doc.paragraphs))
        if suffix == '.md' and role in TEACHERS:
            return bool(path.read_text(encoding='utf-8-sig').strip())
    if role in (*TEACHERS, 'worksheet', 'blackboard') and suffix == '.pdf':
        from pypdf import PdfReader
        return bool(PdfReader(path).pages)
    if role == 'blackboard':
        if suffix == '.pptx':
            from pptx import Presentation
            return any(slide.shapes for slide in Presentation(path).slides)
        if suffix in ('.png', '.jpg', '.jpeg', '.webp'):
            from PIL import Image
            with Image.open(path) as image:
                image.verify()
            return True
    if role == 'interactive' and suffix == '.html':
        text = path.read_text(encoding='utf-8-sig').lower()
        return '<html' in text and '<script' in text
    return False


def _promised_interactive(project, data):
    blueprint = data.get('stages', {}).get('blueprint', {})
    if blueprint.get('interactive_ids'):
        from . import workflow
        audit = workflow.Inspection(project, data)
        audit.evidence(blueprint, 'materials/blueprint', approved=True)
        if not audit.issues:
            return True
    # Unadopted canonical drafts cannot enlarge the current agreed delivery scope.
    for relative, key in (('_state/story.json', 'interactive_nodes'), ('_state/pages.json', 'pages')):
        try:
            state.require_approved(project, relative)
        except ValueError:
            continue
        record = state.read_json(state.resolve(project, relative))
        if key == 'interactive_nodes' and record.get(key):
            return True
        if key == 'pages' and any(p.get('interactive_ids') for p in record.get(key, [])):
            return True
    return False


def check(project, data=None):
    """Return required roles, issues and only explicitly named collection artifacts.

    Existing documents/delivery stage role refs remain valid. A materials-manifest
    adds role arrays (multiple formats/dependencies) and evidence-backed waivers.
    No files are registered, copied, or discovered by directory scanning here.
    """
    if data is None:
        from . import workflow
        data = workflow.load(project)
    full = data.get('current_task', {}).get('mode') == 'full_course'
    modules = {str(x).removeprefix('math-courseware-')
               for x in data.get('current_task', {}).get('modules', [])}
    required_roles = list((*TEACHERS, 'worksheet', 'blackboard')) if full else []
    if full and _promised_interactive(project, data):
        required_roles.append('interactive')
    permitted = set(ROLES) if full or 'studio' in modules else set()
    if 'documents' in modules:
        permitted.update((*TEACHERS, 'worksheet', 'blackboard'))
    if 'interactive' in modules:
        permitted.add('interactive')
    if not permitted:
        return {'allowed': True, 'issues': [], 'required_roles': required_roles, 'artifacts': {}}
    issues, artifacts, found, waived = [], {}, set(), set()
    required_sources = {}
    if full:
        pages = state.resolve(project, '_state/pages.json')
        if pages.is_file():
            required_sources['_state/pages.json'] = state.sha256(pages)
        else:
            issues.append('materials: current canonical pages missing')
        deck = data.get('stages', {}).get('editable', {}).get('files', {}).get('pptx')
        if state.load_project(project).get('artifacts', {}).get('editable-base-pptx'):
            try:
                from . import editable_delivery
                current = editable_delivery.current(project)
                required_sources[current['path']] = current['sha256']
                if not deck:
                    raise ValueError('materials: current editable stage PPT required')
                editable_delivery.current(project, reference=deck)
            except (ValueError, TypeError, KeyError) as exc:
                issues.append(str(exc))
        if deck:
            try:
                _file(project, deck, 'materials/current-pptx')
                required_sources[deck['path']] = deck['sha256']
            except ValueError as exc:
                issues.append(str(exc))

    def add_role(role, items, sources, label):
        if role not in permitted:
            return
        try:
            expected = dict(required_sources)
            if not full:
                module = 'interactive' if role == 'interactive' else 'documents'
                expected.update(current_inputs(project, data, module))
            _sources(project, sources, expected, label)
            if not isinstance(items, list) or not items:
                raise ValueError(label + ': nonempty files list required')
            populated = False
            pending = {}
            for number, item in enumerate(items):
                _delivery_path(state.relative_path(project, item['path']))
                path = _file(project, item, label)
                destination = item.get('destination', state.relative_path(project, path))
                if (not isinstance(destination, str) or not destination or
                        Path(destination).is_absolute() or PureWindowsPath(destination).drive or
                        '..' in PureWindowsPath(destination).parts):
                    raise ValueError(label + ': destination escapes delivery package')
                _delivery_path(destination)
                populated = _populated(path, role) or populated
                pending[f'{label}-{number}'] = {**item, 'role': role,
                    'source_versions': sources, 'destination': destination}
            if not populated:
                raise ValueError(label + ': actual ' + role + ' product required; a specification is insufficient')
            found.add(role)
            artifacts.update(pending)
        except (ValueError, OSError, KeyError, TypeError) as exc:
            issues.append(str(exc))
        except Exception as exc:
            issues.append(label + ': unreadable product (' + type(exc).__name__ + ')')

    for stage_name in ('documents', 'delivery'):
        stage = data.get('stages', {}).get(stage_name, {})
        files = stage.get('files', {})
        for role in ROLES:
            if role in files:
                add_role(role, [files[role]], stage.get('source_versions'), stage_name + '/' + role)
        pointer = files.get('materials-manifest')
        if not pointer:
            continue
        label = stage_name + '/materials-manifest'
        try:
            _delivery_path(state.relative_path(project, pointer['path']))
            path = _file(project, pointer, label)
            manifest = state.read_json(path)
            if manifest.get('schema_version') != '1.0' or not isinstance(manifest.get('roles'), dict):
                raise ValueError(label + ': schema_version 1.0 and roles object required')
            sources = manifest.get('source_versions')
            _sources(project, sources, required_sources, label)
            artifacts[label] = {**pointer, 'source_versions': sources}
            unknown = set(manifest['roles']) - set(ROLES)
            if unknown:
                raise ValueError(label + ': unknown roles: ' + ', '.join(sorted(unknown)))
            for role, items in manifest['roles'].items():
                add_role(role, items, sources, label + '/' + role)
            waivers = manifest.get('waivers', {})
            if not isinstance(waivers, dict):
                raise ValueError(label + ': waivers object required')
            for role, waiver in waivers.items():
                if role not in required_roles:
                    continue
                if not isinstance(waiver, dict) or not str(waiver.get('reason', '')).strip():
                    raise ValueError(label + '/' + role + ': actual waiver reason required')
                evidence = state.read_json(_file(project, waiver.get('evidence'), label + '/' + role + '/waiver'))
                if (not str(evidence.get('user_evidence', '')).strip() or
                        role not in evidence.get('waived_roles', [])):
                    raise ValueError(label + '/' + role + ': actual user scope evidence required')
                _sources(project, evidence.get('source_versions'), required_sources, label + '/' + role + '/waiver')
                waived.add(role)
        except (ValueError, OSError, KeyError, TypeError) as exc:
            issues.append(str(exc))
    for role in required_roles:
        if role not in found and role not in waived:
            issues.append('materials: required role missing: ' + role)
    destinations = {}
    for item in artifacts.values():
        destination = str(PureWindowsPath(item.get('destination', item['path']))).casefold()
        if destination in destinations and destinations[destination] != item['sha256']:
            issues.append('materials: conflicting destination: ' + destination)
        destinations[destination] = item['sha256']
    return {'allowed': not issues, 'issues': list(dict.fromkeys(issues)),
            'required_roles': required_roles, 'artifacts': artifacts}
